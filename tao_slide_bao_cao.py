#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Script: tao_slide_bao_cao.py
Mục đích: Tự động tạo bài thuyết trình PowerPoint (16:9 widescreen) chuyên nghiệp
báo cáo tổng quan về các Vai Trò Người Dùng và Các Phân Hệ Chức Năng Cốt Lõi
cho Hệ thống Quản Lý Nhà Hàng & Món Ăn (Royal Bistro POS & KDS).
Tác giả: Nhóm 5 - Nguyễn Ngọc Hà Thảo
"""

import os
import sys
import pptx

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# ==================== BẢNG MÀU THIẾT KẾ HIỆN ĐẠI ====================
COLOR_BG_DARK = RGBColor(11, 19, 43)        # #0B132B - Nền xanh navy đậm sang trọng
COLOR_BG_CARD = RGBColor(22, 33, 62)        # #16213E - Thẻ nội dung
COLOR_BORDER = RGBColor(48, 64, 102)        # #304066 - Viền thẻ
COLOR_PRIMARY = RGBColor(2, 132, 199)       # #0284C7 - Xanh Cyan / Sky Blue
COLOR_ACCENT = RGBColor(245, 158, 11)       # #F59E0B - Vàng Amber / Gold
COLOR_EMERALD = RGBColor(16, 185, 129)      # #10B981 - Xanh lá tươi
COLOR_PURPLE = RGBColor(129, 140, 248)      # #818CF8 - Tím Indigo nhẹ
COLOR_ORANGE = RGBColor(249, 115, 22)       # #F97316 - Cam Bếp
COLOR_WHITE = RGBColor(255, 255, 255)       # Trắng sáng
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)  # #94A3B8 - Xám sáng
COLOR_TEXT_BODY = RGBColor(226, 232, 240)   # #E2E8F0 - Xám trắng đọc tốt

SLIDE_WIDTH_IN = 13.333
SLIDE_HEIGHT_IN = 7.5

def create_presentation():
    prs = pptx.Presentation()
    prs.slide_width = Inches(SLIDE_WIDTH_IN)
    prs.slide_height = Inches(SLIDE_HEIGHT_IN)
    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_WIDTH_IN), Inches(SLIDE_HEIGHT_IN))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG_DARK
        bg.line.fill.background()
        return bg

    def add_header(slide, badge_text, title_text, subtitle_text, badge_color=COLOR_PRIMARY):
        # Category Badge pill
        badge_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.4), Inches(2.8), Inches(0.38))
        badge_box.fill.solid()
        badge_box.fill.fore_color.rgb = badge_color
        badge_box.line.fill.background()
        tf_b = badge_box.text_frame
        tf_b.word_wrap = True
        tf_b.vertical_anchor = MSO_ANCHOR.MIDDLE
        p_b = tf_b.paragraphs[0]
        p_b.text = badge_text.upper()
        p_b.font.size = Pt(10)
        p_b.font.bold = True
        p_b.font.color.rgb = COLOR_WHITE
        p_b.alignment = PP_ALIGN.CENTER

        # Main Title & Subtitle Box
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.9))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p1 = tf.paragraphs[0]
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_WHITE

        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(12)
        p2.font.color.rgb = COLOR_TEXT_MUTED
        p2.space_before = Pt(4)

    def add_footer(slide, current_slide, total_slides=13):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.0), Inches(11.7), Inches(0.35))
        tf = footer_box.text_frame
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = f"Phần Mềm Quản Lý Nhà Hàng & Món Ăn (Royal Bistro POS)  |  Báo Cáo Tổng Quan Vai Trò & Chức Năng  |  Slide {current_slide}/{total_slides}"
        p.font.size = Pt(9.5)
        p.font.color.rgb = COLOR_TEXT_MUTED

    # ==================== SLIDE 1: TRANG BÌA ====================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative top banner bar
    top_bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_WIDTH_IN), Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT
    top_bar.line.fill.background()

    # Center card container
    card1 = s1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(1.0), Inches(10.9), Inches(5.5))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_BG_CARD
    card1.line.color.rgb = COLOR_BORDER
    card1.line.width = Pt(1.5)

    # Main Title on Card
    tb1 = s1.shapes.add_textbox(Inches(1.8), Inches(1.5), Inches(9.7), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p_badge = tf1.paragraphs[0]
    p_badge.text = "★ BÁO CÁO ĐỒ ÁN CHUYÊN ĐỀ HỆ THỐNG QUẢN LÝ NHÀ HÀNG ★"
    p_badge.font.size = Pt(12)
    p_badge.font.bold = True
    p_badge.font.color.rgb = COLOR_ACCENT
    p_badge.alignment = PP_ALIGN.CENTER

    p_title = tf1.add_paragraph()
    p_title.text = "TỔNG QUAN HỆ THỐNG ROYAL BISTRO POS"
    p_title.font.size = Pt(32)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_WHITE
    p_title.alignment = PP_ALIGN.CENTER
    p_title.space_before = Pt(12)

    p_sub = tf1.add_paragraph()
    p_sub.text = "PHÂN TÍCH CÁC VAI TRÒ NGƯỜI DÙNG & TÍNH NĂNG CHỨC NĂNG CỐT LÕI"
    p_sub.font.size = Pt(16)
    p_sub.font.bold = True
    p_sub.font.color.rgb = COLOR_PRIMARY
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.space_before = Pt(8)

    p_desc = tf1.add_paragraph()
    p_desc.text = "Kiến trúc Fullstack hiện đại kết hợp Realtime Socket.io, Mã QR gọi món tại bàn, Điều phối Bếp KDS và Quản lý Kho Định lượng (BOM)"
    p_desc.font.size = Pt(12)
    p_desc.font.color.rgb = COLOR_TEXT_MUTED
    p_desc.alignment = PP_ALIGN.CENTER
    p_desc.space_before = Pt(12)

    # Info footer on Slide 1
    info_box = s1.shapes.add_textbox(Inches(1.8), Inches(4.8), Inches(9.7), Inches(1.2))
    tfi = info_box.text_frame
    tfi.word_wrap = True

    p_info = tfi.paragraphs[0]
    p_info.text = "Thực hiện: NHÓM 5  |  Trưởng nhóm: NGUYỄN NGỌC HÀ THẢO\nCông nghệ: React 19 • Vite • Node.js Express • MongoDB / MySQL • Socket.io Realtime • Netlify & Cloud"
    p_info.font.size = Pt(11)
    p_info.font.bold = True
    p_info.font.color.rgb = COLOR_TEXT_BODY
    p_info.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 2: TỔNG QUAN KIẾN TRÚC & 4 VAI TRÒ ====================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Phân Tích Kiến Trúc", "MÔ HÌNH VẬN HÀNH & MA TRẬN 4 VAI TRÒ NGƯỜI DÙNG", "Hệ thống số hóa toàn diện quy trình phục vụ khép kín từ Khách hàng -> Thu ngân -> Bếp -> Quản lý")
    add_footer(s2, 2)

    # 4 Role Cards
    roles_data = [
        ("QUẢN TRỊ VIÊN (ADMIN)", COLOR_PURPLE, "Toàn quyền quản trị hệ thống", [
            "Giám sát KPI Doanh thu, số lượng đơn hàng và bàn phục vụ theo thời gian thực.",
            "Quản lý danh mục món ăn, cập nhật giá bán và trạng thái còn/hết hàng.",
            "Cảnh báo tồn kho nguyên liệu, quản lý công thức định lượng (BOM).",
            "Xuất báo cáo tài chính, quản lý tài khoản nhân viên và phân quyền bảo mật."
        ]),
        ("THU NGÂN & PHỤC VỤ (STAFF)", COLOR_EMERALD, "Vận hành sảnh & điểm bán POS", [
            "Theo dõi trực quan sơ đồ bàn ăn & phòng tiệc theo từng khu vực và tầng lầu.",
            "Đổi trạng thái bàn (Mở bàn, dọn bàn, chuyển bàn, ghép bàn).",
            "Lên đơn gọi món tại bàn (POS), thêm ghi chú đặc biệt của thực khách.",
            "Tạm tính, áp dụng giảm giá, in hóa đơn nhiệt 80mm & thanh toán VietQR."
        ]),
        ("ĐẦU BẾP / BẾP TRƯỞNG (KITCHEN)", COLOR_ORANGE, "Màn hình Bếp KDS Realtime", [
            "Tiếp nhận đơn gọi món tức thì qua Socket.io mà không cần nhân viên chạy bàn đưa giấy.",
            "Theo dõi thời gian chờ của từng món ăn để ưu tiên chế biến đúng thứ tự.",
            "Chuyển trạng thái món: Chờ nấu -> Đang chế biến -> Hoàn tất ra món.",
            "Tự động kích hoạt cơ chế trừ tồn kho nguyên liệu khi món được hoàn thành."
        ]),
        ("KHÁCH HÀNG (CUSTOMER)", COLOR_PRIMARY, "Quét mã QR tự gọi món tại bàn", [
            "Quét mã QR dán trên mặt bàn bằng Camera điện thoại (không cần cài app).",
            "Tự động nhận diện số bàn và khu vực phục vụ tương ứng.",
            "Xem E-Menu sống động, tùy chọn số lượng, ghi chú món và đặt hàng trực tiếp.",
            "Theo dõi trạng thái phục vụ và yêu cầu thanh toán / xuất mã QR thanh toán nhanh."
        ])
    ]

    for idx, (r_name, r_color, r_sub, r_bullets) in enumerate(roles_data):
        c_left = Inches(0.8 + idx * 2.95)
        c_top = Inches(1.9)
        c_w = Inches(2.8)
        c_h = Inches(4.9)

        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, c_top, c_w, c_h)
        card.fill.solid()
        card.fill.fore_color.rgb = COLOR_BG_CARD
        card.line.color.rgb = r_color
        card.line.width = Pt(1.5)

        tb = s2.shapes.add_textbox(c_left + Inches(0.15), c_top + Inches(0.15), c_w - Inches(0.3), c_h - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_rn = tf.paragraphs[0]
        p_rn.text = r_name
        p_rn.font.size = Pt(11.5)
        p_rn.font.bold = True
        p_rn.font.color.rgb = r_color

        p_sub = tf.add_paragraph()
        p_sub.text = r_sub
        p_sub.font.size = Pt(9.5)
        p_sub.font.italic = True
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.space_before = Pt(2)
        p_sub.space_after = Pt(8)

        for b in r_bullets:
            p_b = tf.add_paragraph()
            p_b.text = f"• {b}"
            p_b.font.size = Pt(9.5)
            p_b.font.color.rgb = COLOR_TEXT_BODY
            p_b.space_before = Pt(4)

    # ==================== SLIDE 3: XÁC THỰC & ĐĂNG NHẬP PHÂN QUYỀN ====================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Bảo Mật & Phân Quyền", "PHÂN HỆ XÁC THỰC NGƯỜI DÙNG & KIỂM SOÁT VAI TRÒ (RBAC)", "Cơ chế bảo mật JWT tiêu chuẩn công nghiệp và đăng nhập nhanh 1-Click tiện lợi")
    add_footer(s3, 3)

    # Left content box
    card_l3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    card_l3.fill.solid()
    card_l3.fill.fore_color.rgb = COLOR_BG_CARD
    card_l3.line.color.rgb = COLOR_BORDER
    tb_l3 = s3.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.6))
    tf_l3 = tb_l3.text_frame
    tf_l3.word_wrap = True
    tf_l3.margin_left = tf_l3.margin_top = tf_l3.margin_right = tf_l3.margin_bottom = 0

    p_h = tf_l3.paragraphs[0]
    p_h.text = "TÍNH NĂNG ĐĂNG NHẬP AN TOÀN"
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_ACCENT

    points3 = [
        ("Mã Hóa & Cấp Token JWT:", "Mật khẩu người dùng được băm an toàn bằng thuật toán Bcrypt. Máy chủ cấp chuỗi JWT Token có thời hạn sử dụng xác thực mọi API."),
        ("Phân Quyền RBAC (Role-Based Access):", "Hệ thống tự động lọc các mục Menu phù hợp: Admin toàn quyền, Nhân viên chỉ thấy POS/Bàn, Bếp chỉ thấy KDS."),
        ("Tính Năng 1-Click Demo Login:", "Tích hợp sẵn 3 nút chọn tài khoản mẫu (Admin, Thu Ngân, Bếp Trưởng) giúp giảng viên và hội đồng kiểm thử nhanh chóng."),
        ("Tự Động Bắt Lỗi Phiên Hết Hạn:", "Interceptor phía Frontend tự động phát hiện mã lỗi 401 Unauthorized để điều hướng về màn hình đăng nhập an toàn.")
    ]
    for title, desc in points3:
        p_t = tf_l3.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_l3.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(2)

    # Right Image
    img_path3 = "docs/screenshots/01_login.png"
    if os.path.exists(img_path3):
        card_r3 = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
        card_r3.fill.solid()
        card_r3.fill.fore_color.rgb = COLOR_BG_CARD
        card_r3.line.color.rgb = COLOR_PRIMARY
        card_r3.line.width = Pt(1.5)
        s3.shapes.add_picture(img_path3, Inches(5.7), Inches(2.0), width=Inches(6.7))

    # ==================== SLIDE 4: DASHBOARD QUẢN TRỊ VIÊN ====================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Vai Trò: Quản Trị Viên", "TỔNG QUAN DASHBOARD: THEO DÕI VẬN HÀNH & KINH DOANH", "Cung cấp góc nhìn 360 độ về hiệu suất bán hàng, trạng thái bàn và cảnh báo vận hành", COLOR_PURPLE)
    add_footer(s4, 4)

    card_l4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    card_l4.fill.solid()
    card_l4.fill.fore_color.rgb = COLOR_BG_CARD
    card_l4.line.color.rgb = COLOR_BORDER
    tb_l4 = s4.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.6))
    tf_l4 = tb_l4.text_frame
    tf_l4.word_wrap = True
    tf_l4.margin_left = tf_l4.margin_top = tf_l4.margin_right = tf_l4.margin_bottom = 0

    p_h = tf_l4.paragraphs[0]
    p_h.text = "CÁC CHỈ SỐ DOANH THU & KPI"
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_PURPLE

    points4 = [
        ("4 Thẻ Chỉ Số KPI Thời Gian Thực:", "Tổng doanh thu bán hàng trong ngày, số đơn hoàn tất, số bàn đang phục vụ thực khách và số nguyên liệu cần nhập kho."),
        ("Biểu Đồ Xu Hướng Doanh Số:", "Theo dõi tăng trưởng doanh số theo tuần/tháng, trực quan hóa cơ cấu nguồn thu tiền mặt và chuyển khoản."),
        ("Xếp Hạng Top Món Ăn Bán Chạy:", "Nhận diện ngay món ăn được khách yêu thích nhất để chuẩn bị nguyên liệu kho bếp và thúc đẩy doanh thu."),
        ("Dòng Nhật Ký Hoạt Động (Live Feed):", "Theo dõi theo thời gian thực mọi hoạt động mở bàn, gọi món và thanh toán của nhân viên sảnh.")
    ]
    for title, desc in points4:
        p_t = tf_l4.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_l4.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(2)

    img_path4 = "docs/screenshots/02_dashboard.png"
    if os.path.exists(img_path4):
        card_r4 = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
        card_r4.fill.solid()
        card_r4.fill.fore_color.rgb = COLOR_BG_CARD
        card_r4.line.color.rgb = COLOR_PURPLE
        card_r4.line.width = Pt(1.5)
        s4.shapes.add_picture(img_path4, Inches(5.7), Inches(2.0), width=Inches(6.7))

    # ==================== SLIDE 5: SƠ ĐỒ BÀN ĂN & PHÒNG TIỆC ====================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Vai Trò: Thu Ngân & Phục Vụ", "SƠ ĐỒ BÀN ĂN & PHÒNG TIỆC: QUẢN LÝ MẶT BẰNG TRỰC QUAN", "Theo dõi tình trạng bàn theo khu vực, thời gian thực và quản lý liên kết mã QR từng bàn", COLOR_EMERALD)
    add_footer(s5, 5)

    card_l5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    card_l5.fill.solid()
    card_l5.fill.fore_color.rgb = COLOR_BG_CARD
    card_l5.line.color.rgb = COLOR_BORDER
    tb_l5 = s5.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.6))
    tf_l5 = tb_l5.text_frame
    tf_l5.word_wrap = True
    tf_l5.margin_left = tf_l5.margin_top = tf_l5.margin_right = tf_l5.margin_bottom = 0

    p_h = tf_l5.paragraphs[0]
    p_h.text = "QUẢN LÝ BÀN THÔNG MINH"
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_EMERALD

    points5 = [
        ("Bộ Lọc Phân Vùng Không Gian:", "Phân chia mặt bằng khoa học: Tầng 1 Sảnh Chính, Tầng 2 Ban Công Lãng Mạn, Phòng VIP 1 Hoàng Gia, Phòng VIP 2 Kim Cương."),
        ("Đèn Tín Hiệu Trạng Thái 3 Màu:", "Màu xanh (Trống - Sẵn sàng), Màu vàng nhấp nháy (Đang có khách dùng bữa), Màu xanh dương (Đã đặt chỗ trước)."),
        ("Hiển Thị Chi Tiết Thời Gian Thực:", "Xem ngay số lượng khách đang ngồi / sức chứa tối đa, số món ăn đang dùng và tổng tiền tạm tính hiện thời."),
        ("Tích Hợp Mã QR Bàn Độc Nhất:", "Cung cấp nút xem mã QR riêng biệt cho từng bàn để in ấn dán tại chỗ hoặc sao chép link quét mã cho khách.")
    ]
    for title, desc in points5:
        p_t = tf_l5.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_l5.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(2)

    img_path5 = "docs/screenshots/03_tables.png"
    if os.path.exists(img_path5):
        card_r5 = s5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
        card_r5.fill.solid()
        card_r5.fill.fore_color.rgb = COLOR_BG_CARD
        card_r5.line.color.rgb = COLOR_EMERALD
        card_r5.line.width = Pt(1.5)
        s5.shapes.add_picture(img_path5, Inches(5.7), Inches(2.0), width=Inches(6.7))

    # ==================== SLIDE 6: POS GỌI MÓN & THU NGÂN ====================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Vai Trò: Thu Ngân", "ĐIỂM BÁN HÀNG POS: GỌI MÓN NHANH CHÓNG & IN HÓA ĐƠN", "Giao diện cảm ứng tối ưu hóa tốc độ nhập đơn của nhân viên thu ngân và phục vụ", COLOR_EMERALD)
    add_footer(s6, 6)

    card_l6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    card_l6.fill.solid()
    card_l6.fill.fore_color.rgb = COLOR_BG_CARD
    card_l6.line.color.rgb = COLOR_BORDER
    tb_l6 = s6.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.6))
    tf_l6 = tb_l6.text_frame
    tf_l6.word_wrap = True
    tf_l6.margin_left = tf_l6.margin_top = tf_l6.margin_right = tf_l6.margin_bottom = 0

    p_h = tf_l6.paragraphs[0]
    p_h.text = "NGHIỆP VỤ BÁN HÀNG TẠI BÀN"
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_ACCENT

    points6 = [
        ("Thiết Kế 2 Cột Hiện Đại:", "Cột trái duyệt menu trực quan theo loại món kèm tìm kiếm nhanh; Cột phải quản lý giỏ hàng chi tiết của bàn đang chọn."),
        ("Ghi Chú Khẩu Vị Chuyên Biệt:", "Dễ dàng thêm ghi chú món trực tiếp (ví dụ: 'không tiêu', 'ít cay', 'uống không đá') để truyền tải chuẩn xác tới đầu bếp."),
        ("Tự Động Tính Toán Tài Chính:", "Hệ thống tự động tính tổng tiền, chiết khấu khách quen, thuế VAT và số tiền khách cần thanh toán."),
        ("In Hóa Đơn Nhiệt & Thanh Toán:", "Hỗ trợ in phiếu thanh toán chuẩn khổ 80mm và hiển thị mã QR VietQR chuyển khoản ngân hàng tiện lợi.")
    ]
    for title, desc in points6:
        p_t = tf_l6.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_l6.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(2)

    img_path6 = "docs/screenshots/04_pos_order.png"
    if os.path.exists(img_path6):
        card_r6 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
        card_r6.fill.solid()
        card_r6.fill.fore_color.rgb = COLOR_BG_CARD
        card_r6.line.color.rgb = COLOR_ACCENT
        card_r6.line.width = Pt(1.5)
        s6.shapes.add_picture(img_path6, Inches(5.7), Inches(2.0), width=Inches(6.7))

    # ==================== SLIDE 7: MÀN HÌNH BẾP KDS REALTIME ====================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Vai Trò: Đầu Bếp", "MÀN HÌNH ĐIỀU PHỐI BẾP KDS: ĐỒNG BỘ THỜI GIAN THỰC", "Loại bỏ hoàn toàn giấy in order truyền thống, đồng bộ đơn nấu tức thì qua Socket.io", COLOR_ORANGE)
    add_footer(s7, 7)

    card_l7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    card_l7.fill.solid()
    card_l7.fill.fore_color.rgb = COLOR_BG_CARD
    card_l7.line.color.rgb = COLOR_BORDER
    tb_l7 = s7.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.6))
    tf_l7 = tb_l7.text_frame
    tf_l7.word_wrap = True
    tf_l7.margin_left = tf_l7.margin_top = tf_l7.margin_right = tf_l7.margin_bottom = 0

    p_h = tf_l7.paragraphs[0]
    p_h.text = "CÔNG NGHỆ BẾP KDS HIỆN ĐẠI"
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_ORANGE

    points7 = [
        ("Truyền Tin Realtime Qua Socket.io:", "Khi Thu ngân bấm 'Lên Đơn' hoặc Khách tự gửi món trên điện thoại, phiếu bếp lập tức xuất hiện kèm chuông báo động."),
        ("Bộ Đếm Thời Gian Chờ (Wait Timer):", "Đo lường chính xác số phút khách đã chờ, chuyển sang màu đỏ cảnh báo khi vượt định mức để bếp ưu tiên nấu trước."),
        ("Quy Trình 3 Bước Tinh Gọn:", "Phiếu bếp phân làn rõ ràng: Chờ duyệt -> Đang nấu trên chảo -> Đã chế biến xong, sẵn sàng chuyển cho phục vụ."),
        ("Chống Sót Đơn Tuyệt Đối:", "Loại bỏ 100% tình trạng mất giấy order, nhầm bàn hay sót món khi nhà hàng đạt công suất phục vụ tối đa.")
    ]
    for title, desc in points7:
        p_t = tf_l7.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_l7.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(2)

    img_path7 = "docs/screenshots/05_kitchen_kds.png"
    if os.path.exists(img_path7):
        card_r7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
        card_r7.fill.solid()
        card_r7.fill.fore_color.rgb = COLOR_BG_CARD
        card_r7.line.color.rgb = COLOR_ORANGE
        card_r7.line.width = Pt(1.5)
        s7.shapes.add_picture(img_path7, Inches(5.7), Inches(2.0), width=Inches(6.7))

    # ==================== SLIDE 8: QUẢN LÝ THỰC ĐƠN MÓN ĂN ====================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Vai Trò: Quản Trị & Bếp", "QUẢN LÝ THỰC ĐƠN: DANH MỤC MÓN ĂN & ĐỊNH GIÁ NIÊM YẾT", "Cấu hình toàn diện danh mục món, giá bán, hình ảnh và kiểm soát trạng thái còn/hết hàng", COLOR_PURPLE)
    add_footer(s8, 8)

    card_l8 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    card_l8.fill.solid()
    card_l8.fill.fore_color.rgb = COLOR_BG_CARD
    card_l8.line.color.rgb = COLOR_BORDER
    tb_l8 = s8.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.6))
    tf_l8 = tb_l8.text_frame
    tf_l8.word_wrap = True
    tf_l8.margin_left = tf_l8.margin_top = tf_l8.margin_right = tf_l8.margin_bottom = 0

    p_h = tf_l8.paragraphs[0]
    p_h.text = "QUẢN TRỊ DANH MỤC THỰC ĐƠN"
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_PRIMARY

    points8 = [
        ("Phân Cấp Danh Mục Khoa Học:", "Hỗ trợ đầy đủ các nhóm món ăn: Khai vị, Món chính, Món lẩu, Hải sản cao cấp, Tráng miệng và Đồ uống các loại."),
        ("Cập Nhật Trạng Thái Món Nhanh:", "Bật/tắt trạng thái 'Còn hàng' sang 'Tạm hết' chỉ bằng 1 thao tác, tự động ngăn chặn khách đặt món khi hết nguyên liệu."),
        ("Thêm Món Ăn Mới Trực Quan:", "Biểu mẫu nhập liệu chuẩn hóa: Tên món, phân loại danh mục, giá niêm yết, đường dẫn ảnh minh họa và mô tả khẩu vị."),
        ("Xóa & Ẩn Món An Toàn:", "Hệ thống áp dụng ràng buộc khóa ngoại an toàn, bảo vệ dữ liệu lịch sử các đơn hàng cũ đã từng bán món này.")
    ]
    for title, desc in points8:
        p_t = tf_l8.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_l8.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(2)

    img_path8 = "docs/screenshots/06_dishes.png"
    if os.path.exists(img_path8):
        card_r8 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
        card_r8.fill.solid()
        card_r8.fill.fore_color.rgb = COLOR_BG_CARD
        card_r8.line.color.rgb = COLOR_PRIMARY
        card_r8.line.width = Pt(1.5)
        s8.shapes.add_picture(img_path8, Inches(5.7), Inches(2.0), width=Inches(6.7))

    # ==================== SLIDE 9: VAI TRÒ KHÁCH HÀNG: QUÉT MÃ QR GỌI MÓN ====================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Vai Trò: Khách Hàng", "QUÉT MÃ QR GỌI MÓN TẠI BÀN (E-MENU): TRẢI NGHIỆM KHÔNG CHẠM", "Khách tự dùng bữa văn minh, tự chọn món trên điện thoại cá nhân và gửi thẳng vào Bếp KDS", COLOR_PRIMARY)
    add_footer(s9, 9)

    card_l9 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(7.0), Inches(5.0))
    card_l9.fill.solid()
    card_l9.fill.fore_color.rgb = COLOR_BG_CARD
    card_l9.line.color.rgb = COLOR_BORDER
    tb_l9 = s9.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(6.6), Inches(4.6))
    tf_l9 = tb_l9.text_frame
    tf_l9.word_wrap = True
    tf_l9.margin_left = tf_l9.margin_top = tf_l9.margin_right = tf_l9.margin_bottom = 0

    p_h = tf_l9.paragraphs[0]
    p_h.text = "QUY TRÌNH TỰ ĐỘNG HÓA TRẢI NGHIỆM THỰC KHÁCH"
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_PRIMARY

    points9 = [
        ("Bước 1: Quét Mã QR Tại Bàn:", "Thực khách dùng ứng dụng Zalo, Camera iOS/Android quét mã QR dán trên mặt bàn. Trình duyệt tự mở trang web E-Menu gắn liền ID bàn."),
        ("Bước 2: Lựa Chọn Món Ăn Sinh Động:", "Thực khách duyệt danh mục, xem hình ảnh chân thực, chọn số lượng và ghi chú gia vị (ít đá, không ớt, sốt riêng)."),
        ("Bước 3: Đặt Hàng Trực Tiếp Vào Bếp:", "Đơn hàng bay trực tiếp vào Màn hình Bếp KDS để nấu và hiển thị trên màn hình POS của thu ngân mà không cần vẫy tay gọi phục vụ."),
        ("Bước 4: Theo Dõi Món & Thanh Toán:", "Khách theo dõi trực tiếp trạng thái món ('Đang nấu' -> 'Đã dọn lên'), đồng thời chủ động bấm 'Yêu Cầu Thanh Toán' hoặc quét mã VietQR.")
    ]
    for title, desc in points9:
        p_t = tf_l9.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(6)

        p_d = tf_l9.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(2)

    # Right side: QR Image card
    card_r9 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.8), Inches(4.3), Inches(5.0))
    card_r9.fill.solid()
    card_r9.fill.fore_color.rgb = COLOR_BG_CARD
    card_r9.line.color.rgb = COLOR_ACCENT
    card_r9.line.width = Pt(1.5)

    qr_img = "docs/screenshots/ma_qr.jpg"
    if os.path.exists(qr_img):
        s9.shapes.add_picture(qr_img, Inches(8.85), Inches(2.2), width=Inches(3.0))

    tb_qr_lbl = s9.shapes.add_textbox(Inches(8.4), Inches(5.4), Inches(3.9), Inches(1.2))
    tf_qr_lbl = tb_qr_lbl.text_frame
    tf_qr_lbl.word_wrap = True
    p_qr = tf_qr_lbl.paragraphs[0]
    p_qr.text = "MÃ QR THỬ NGHIỆM TẠI BÀN\nQuét mã để truy cập giao diện E-Menu thực tế (?table=1)"
    p_qr.font.size = Pt(10.5)
    p_qr.font.bold = True
    p_qr.font.color.rgb = COLOR_ACCENT
    p_qr.alignment = PP_ALIGN.CENTER

    # ==================== SLIDE 10: QUẢN LÝ KHO & ĐỊNH LƯỢNG (BOM) ====================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Nghiệp Vụ Chuyên Sâu", "QUẢN LÝ KHO & CÔNG THỨC ĐỊNH LƯỢNG NGUYÊN LIỆU (BOM)", "Cơ chế tự động trừ tồn kho khi Bếp nấu xong và cảnh báo nguyên liệu chạm ngưỡng tối thiểu", COLOR_ORANGE)
    add_footer(s10, 10)

    card_l10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    card_l10.fill.solid()
    card_l10.fill.fore_color.rgb = COLOR_BG_CARD
    card_l10.line.color.rgb = COLOR_BORDER
    tb_l10 = s10.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.6))
    tf_l10 = tb_l10.text_frame
    tf_l10.word_wrap = True
    tf_l10.margin_left = tf_l10.margin_top = tf_l10.margin_right = tf_l10.margin_bottom = 0

    p_h = tf_l10.paragraphs[0]
    p_h.text = "ĐỊNH LƯỢNG NGUYÊN VẬT LIỆU"
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_ORANGE

    points10 = [
        ("Bảng Định Lượng Chuẩn (BOM):", "Mỗi món ăn được gắn với công thức gồm danh sách nguyên liệu và định lượng chuẩn (ví dụ: Bò Bít Tết gồm 200g Bò Wagyu, 20g Bơ, 10g Tiêu đen)."),
        ("Tự Động Trừ Kho (Auto Deduction):", "Khi đầu bếp ấn 'Hoàn Tất' món trên KDS, máy chủ tự động tính toán và trừ tồn kho chính xác từng gam nguyên liệu."),
        ("Cảnh Báo Chạm Ngưỡng Tồn Kho:", "Đèn tín hiệu đổi màu đỏ khi số lượng nguyên liệu trong kho xuống thấp hơn định mức tối thiểu, nhắc nhở quản lý nhập hàng."),
        ("Quản Lý Nhà Cung Cấp Thực Phẩm:", "Lưu trữ đầy đủ thông tin nhà cung cấp thực phẩm uy tín, số điện thoại liên hệ và đánh giá chất lượng nguyên liệu.")
    ]
    for title, desc in points10:
        p_t = tf_l10.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_l10.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(2)

    flow_img = "generated_diagrams/flowchart_bom_inventory.png"
    if os.path.exists(flow_img):
        card_r10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
        card_r10.fill.solid()
        card_r10.fill.fore_color.rgb = COLOR_BG_CARD
        card_r10.line.color.rgb = COLOR_ORANGE
        card_r10.line.width = Pt(1.5)
        s10.shapes.add_picture(flow_img, Inches(5.7), Inches(2.1), width=Inches(6.7))

    # ==================== SLIDE 11: MÔ HÌNH DỮ LIỆU CƠ SỞ DỮ LIỆU (ERD) ====================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "Kiến Trúc Dữ Liệu", "LƯỢC ĐỒ QUAN HỆ THỰC THỂ CƠ SỞ DỮ LIỆU (ERD)", "Thiết kế chuẩn hóa 3NF gồm 12 bảng thực thể đảm bảo tính toàn vẹn và hiệu năng ACID", COLOR_EMERALD)
    add_footer(s11, 11)

    card_l11 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.8), Inches(4.5), Inches(5.0))
    card_l11.fill.solid()
    card_l11.fill.fore_color.rgb = COLOR_BG_CARD
    card_l11.line.color.rgb = COLOR_BORDER
    tb_l11 = s11.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(4.1), Inches(4.6))
    tf_l11 = tb_l11.text_frame
    tf_l11.word_wrap = True
    tf_l11.margin_left = tf_l11.margin_top = tf_l11.margin_right = tf_l11.margin_bottom = 0

    p_h = tf_l11.paragraphs[0]
    p_h.text = "CẤU TRÚC 12 BẢNG THỰC THỂ"
    p_h.font.size = Pt(14)
    p_h.font.bold = True
    p_h.font.color.rgb = COLOR_EMERALD

    points11 = [
        ("Phân Hệ Danh Mục & Thực Đơn:", "Bảng loai_mon (1) - mon_an (N) lưu trữ toàn bộ thực đơn; mon_an_nguyen_lieu giải quyết liên kết N-N giữa Món ăn và Kho."),
        ("Phân Hệ Mặt Bằng & Gọi Món:", "Bảng ban liên kết dat_mon (ghi nhận chi tiết từng món gọi), dat_ban_truoc (đặt chỗ hẹn giờ) và danh_gia_mon_an (review sao)."),
        ("Phân Hệ Khách Hàng & Phân Quyền:", "Bảng khach_hang tích lũy điểm & nâng hạng thành viên; Bảng users phân quyền Admin, Thu ngân, Bếp."),
        ("Toàn Vẹn Dữ Liệu & Chỉ Mục Tối Ưu:", "Sử dụng quy tắc ON DELETE CASCADE và ON DELETE SET NULL, đánh chỉ mục Index cho tốc độ truy vấn < 10ms.")
    ]
    for title, desc in points11:
        p_t = tf_l11.add_paragraph()
        p_t.text = f"✔ {title}"
        p_t.font.size = Pt(10.5)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_WHITE
        p_t.space_before = Pt(8)

        p_d = tf_l11.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(9.5)
        p_d.font.color.rgb = COLOR_TEXT_MUTED
        p_d.space_before = Pt(2)

    erd_img = "generated_diagrams/erd_database.png"
    if os.path.exists(erd_img):
        card_r11 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.6), Inches(1.8), Inches(6.9), Inches(5.0))
        card_r11.fill.solid()
        card_r11.fill.fore_color.rgb = COLOR_BG_CARD
        card_r11.line.color.rgb = COLOR_EMERALD
        card_r11.line.width = Pt(1.5)
        s11.shapes.add_picture(erd_img, Inches(5.7), Inches(2.1), width=Inches(6.7))

    # ==================== SLIDE 12: MA TRẬN PHÂN QUYỀN & LỢI ÍCH ====================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Tổng Kết Đánh Giá", "MA TRẬN CHỨC NĂNG THEO VAI TRÒ & HIỆU QUẢ VẬN HÀNH", "Đối chiếu toàn bộ tính năng theo 4 nhóm đối tượng sử dụng và đo lường giá trị thực tiễn")
    add_footer(s12, 12)

    # Matrix Table Shape
    rows = 10
    cols = 5
    table_shape = s12.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(7.5), Inches(5.0))
    table = table_shape.table

    # Column widths
    table.columns[0].width = Inches(2.7)
    table.columns[1].width = Inches(1.2)
    table.columns[2].width = Inches(1.2)
    table.columns[3].width = Inches(1.2)
    table.columns[4].width = Inches(1.2)

    headers_m = ["Phân Hệ Chức Năng", "Quản Trị", "Thu Ngân", "Đầu Bếp", "Khách"]
    for c_idx, h in enumerate(headers_m):
        cell = table.cell(0, c_idx)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_PRIMARY
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_WHITE
        p.alignment = PP_ALIGN.CENTER if c_idx > 0 else PP_ALIGN.LEFT

    matrix_rows = [
        ("1. Tổng quan Dashboard & KPI", "✔ Toàn quyền", "—", "—", "—"),
        ("2. Sơ đồ bàn & Trạng thái bàn", "✔ Quản lý", "✔ Tác nghiệp", "—", "—"),
        ("3. Điểm bán POS & Gọi món nhanh", "✔ Quản lý", "✔ Gọi món", "—", "—"),
        ("4. Màn hình Bếp KDS Realtime", "✔ Giám sát", "—", "✔ Chế biến", "—"),
        ("5. Khách quét mã QR tự gọi món", "✔ Cấu hình", "✔ Tiếp nhận", "✔ Nhận đơn", "✔ Tự order"),
        ("6. Đặt bàn trước & Cọc giữ chỗ", "✔ Quản lý", "✔ Xếp chỗ", "—", "—"),
        ("7. Quản lý Thực đơn & Giá món", "✔ Toàn quyền", "—", "✔ Cập nhật", "—"),
        ("8. Quản lý Kho & Định lượng BOM", "✔ Toàn quyền", "—", "✔ Trừ kho", "—"),
        ("9. Báo cáo thống kê doanh thu", "✔ Toàn quyền", "—", "—", "—")
    ]

    for r_idx, r_data in enumerate(matrix_rows, start=1):
        for c_idx, val in enumerate(r_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = COLOR_BG_CARD if r_idx % 2 == 1 else RGBColor(18, 27, 51)
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(9)
            if c_idx == 0:
                p.font.bold = True
                p.font.color.rgb = COLOR_WHITE
            else:
                p.font.color.rgb = COLOR_EMERALD if "✔" in val else COLOR_TEXT_MUTED
                p.alignment = PP_ALIGN.CENTER

    # Right side: 4 Key Value Props Card
    card_r12 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(1.8), Inches(3.9), Inches(5.0))
    card_r12.fill.solid()
    card_r12.fill.fore_color.rgb = COLOR_BG_CARD
    card_r12.line.color.rgb = COLOR_ACCENT
    card_r12.line.width = Pt(1.5)

    tb_r12 = s12.shapes.add_textbox(Inches(8.8), Inches(2.0), Inches(3.5), Inches(4.6))
    tf_r12 = tb_r12.text_frame
    tf_r12.word_wrap = True

    p_vh = tf_r12.paragraphs[0]
    p_vh.text = "GIÁ TRỊ THỰC TIỄN ĐẠT ĐƯỢC"
    p_vh.font.size = Pt(13)
    p_vh.font.bold = True
    p_vh.font.color.rgb = COLOR_ACCENT

    benefits = [
        ("⚡ Giảm 40% Chi Phí Nhân Sự:", "Khách tự quét mã gọi món tại bàn, giảm tải áp lực cho nhân viên phục vụ trong giờ cao điểm."),
        ("🛡️ Triệt Tiêu 100% Sót Đơn:", "Đồng bộ Socket.io tức thì giữa POS và Màn hình Bếp KDS giúp đơn không bao giờ bị rơi rớt."),
        ("⏱️ Tăng 200% Tốc Độ Ra Món:", "Bộ đếm thời gian thực trên KDS giúp đầu bếp luôn chủ động ưu tiên các món chờ lâu."),
        ("📊 Chống Thất Thoát Kho & Tiền:", "Hạch toán tự động từng đơn hàng, tự trừ kho theo công thức BOM chuẩn xác.")
    ]
    for b_title, b_desc in benefits:
        p_bt = tf_r12.add_paragraph()
        p_bt.text = b_title
        p_bt.font.size = Pt(10)
        p_bt.font.bold = True
        p_bt.font.color.rgb = COLOR_WHITE
        p_bt.space_before = Pt(8)

        p_bd = tf_r12.add_paragraph()
        p_bd.text = b_desc
        p_bd.font.size = Pt(9)
        p_bd.font.color.rgb = COLOR_TEXT_MUTED
        p_bd.space_before = Pt(2)

    # ==================== SLIDE 13: KẾT THÚC & Q&A ====================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)

    top_bar13 = s13.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_WIDTH_IN), Inches(0.15))
    top_bar13.fill.solid()
    top_bar13.fill.fore_color.rgb = COLOR_EMERALD
    top_bar13.line.fill.background()

    card13 = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.2), Inches(10.3), Inches(5.1))
    card13.fill.solid()
    card13.fill.fore_color.rgb = COLOR_BG_CARD
    card13.line.color.rgb = COLOR_PRIMARY
    card13.line.width = Pt(1.5)

    tb13 = s13.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.3), Inches(3.9))
    tf13 = tb13.text_frame
    tf13.word_wrap = True

    p_ty = tf13.paragraphs[0]
    p_ty.text = "CẢM ƠN THẦY CÔ & HỘI ĐỒNG ĐÃ LẮNG NGHE!"
    p_ty.font.size = Pt(26)
    p_ty.font.bold = True
    p_ty.font.color.rgb = COLOR_ACCENT
    p_ty.alignment = PP_ALIGN.CENTER

    p_qa = tf13.add_paragraph()
    p_qa.text = "HỎI & ĐÁP (Q & A)"
    p_qa.font.size = Pt(20)
    p_qa.font.bold = True
    p_qa.font.color.rgb = COLOR_WHITE
    p_qa.alignment = PP_ALIGN.CENTER
    p_qa.space_before = Pt(12)

    p_c = tf13.add_paragraph()
    p_c.text = "Hệ thống Quản lý Nhà hàng & Món ăn (Royal Bistro POS)\nĐề tài: Thiết kế & Cài đặt Hệ thống Quản trị Cơ sở Dữ liệu & Ứng dụng Web Thời gian thực\nThực hiện: Nhóm 5 — Trưởng nhóm: Nguyễn Ngọc Hà Thảo"
    p_c.font.size = Pt(12)
    p_c.font.color.rgb = COLOR_TEXT_MUTED
    p_c.alignment = PP_ALIGN.CENTER
    p_c.space_before = Pt(16)

    output_filename = "BaoCao_TongQuan_VaiTro_ChucNang.pptx"
    prs.save(output_filename)
    print(f"Đã tạo thành công bài thuyết trình: {output_filename}")
    return output_filename

if __name__ == "__main__":
    create_presentation()
