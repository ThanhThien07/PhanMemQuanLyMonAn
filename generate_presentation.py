import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- COLOR PALETTE (BRIGHT, MODERN & HIGH CONTRAST) ---
    LIGHT_BG = RGBColor(248, 250, 252)    # Slate 50
    SLIDE_WHITE = RGBColor(255, 255, 255)
    PRIMARY = RGBColor(30, 58, 138)       # Royal Blue 900
    VIBRANT_BLUE = RGBColor(37, 99, 235)  # Blue 600
    SKY_BLUE = RGBColor(2, 132, 199)      # Sky 600
    GOLD = RGBColor(217, 119, 6)          # Amber 600
    EMERALD = RGBColor(5, 150, 105)       # Emerald 600
    RED = RGBColor(220, 38, 38)           # Red 600
    TEXT_DARK = RGBColor(15, 23, 42)      # Slate 900
    TEXT_MUTED = RGBColor(71, 85, 105)    # Slate 600
    CARD_BG = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
    CARD_ALT = RGBColor(241, 245, 249)    # Slate 100
    WHITE = RGBColor(255, 255, 255)

    media_dir = r'c:\Users\LENOVO\OneDrive\Documents\My Data Sources\PhanMemQuanLyMonAn\docx_media'
    TOTAL_SLIDES = 11

    def set_slide_bg(slide, color=LIGHT_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text="", tag_width=3.2):
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.15))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = WHITE
        header_box.line.color.rgb = CARD_BORDER
        header_box.line.width = Pt(1)

        accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.06))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = VIBRANT_BLUE
        accent_bar.line.fill.background()

        tag_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.14), Inches(tag_width), Inches(0.26))
        tag_shape.fill.solid()
        tag_shape.fill.fore_color.rgb = RGBColor(238, 242, 255)
        tag_shape.line.color.rgb = RGBColor(199, 210, 254)
        tag_shape.line.width = Pt(0.75)
        tf_tag = tag_shape.text_frame
        tf_tag.word_wrap = False
        tf_tag.margin_top = Inches(0.02)
        tf_tag.margin_bottom = Inches(0.02)
        tf_tag.margin_left = Inches(0.08)
        tf_tag.margin_right = Inches(0.08)
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.alignment = PP_ALIGN.CENTER
        p_tag.font.name = "Segoe UI"
        p_tag.font.size = Pt(8.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = VIBRANT_BLUE

        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.44), Inches(11.7), Inches(0.65))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_left = Inches(0)
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Segoe UI"
        p.font.size = Pt(19)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(10.5)
            p2.font.color.rgb = TEXT_MUTED

    def add_footer(slide, current_idx, total_slides=TOTAL_SLIDES):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(11.733), Inches(0.32))
        tf = footer_box.text_frame
        tf.margin_left = Inches(0)
        tf.margin_top = Inches(0)
        p = tf.paragraphs[0]
        p.text = f"Đồ Án Chuyên Đề Back-End: Royal Bistro | Sinh viên: Nguyễn Hoàng Hưng – 501250384 (K25)                                 Trang {current_idx}/{total_slides}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(9)
        p.font.color.rgb = RGBColor(148, 163, 184)

    # ============================================================
    # SLIDE 1: TRANG BÌA (BRIGHT, MODERN LIGHT THEME - DỄ NHÌN)
    # ============================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide1, LIGHT_BG)

    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = VIBRANT_BLUE
    top_bar.line.fill.background()

    sch_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.55), Inches(11.533), Inches(0.5))
    sch_box.fill.solid()
    sch_box.fill.fore_color.rgb = WHITE
    sch_box.line.color.rgb = CARD_BORDER
    sch_box.line.width = Pt(1)
    sch_tf = sch_box.text_frame
    sch_tf.margin_top = Inches(0.08)
    sch_p = sch_tf.paragraphs[0]
    sch_p.text = "🏛️ TRƯỜNG CAO ĐẲNG CÔNG NGHỆ THÔNG TIN TP.HCM • KHOA CÔNG NGHỆ THÔNG TIN – ĐIỆN TỬ"
    sch_p.alignment = PP_ALIGN.CENTER
    sch_p.font.name = "Segoe UI"
    sch_p.font.size = Pt(11)
    sch_p.font.bold = True
    sch_p.font.color.rgb = PRIMARY

    main_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.25), Inches(11.533), Inches(2.35))
    main_card.fill.solid()
    main_card.fill.fore_color.rgb = WHITE
    main_card.line.color.rgb = RGBColor(199, 210, 254)
    main_card.line.width = Pt(1.5)
    
    m_tf = main_card.text_frame
    m_tf.word_wrap = True
    m_tf.margin_top = Inches(0.22)
    m_tf.margin_left = Inches(0.4)
    m_tf.margin_right = Inches(0.4)

    mp0 = m_tf.paragraphs[0]
    mp0.text = "BÁO CÁO ĐỒ ÁN MÔN HỌC: CHUYÊN ĐỀ BACK-END"
    mp0.alignment = PP_ALIGN.CENTER
    mp0.font.name = "Segoe UI"
    mp0.font.size = Pt(13)
    mp0.font.bold = True
    mp0.font.color.rgb = VIBRANT_BLUE

    mp1 = m_tf.add_paragraph()
    mp1.text = "XÂY DỰNG HỆ THỐNG BACK-END CHO PHẦN MỀM\nQUẢN LÝ NHÀ HÀNG & GỌI MÓN (ROYAL BISTRO)"
    mp1.alignment = PP_ALIGN.CENTER
    mp1.font.name = "Segoe UI"
    mp1.font.size = Pt(25)
    mp1.font.bold = True
    mp1.font.color.rgb = PRIMARY
    mp1.space_before = Pt(6)

    mp2 = m_tf.add_paragraph()
    mp2.text = "Kiến trúc máy chủ RESTful API & WebSocket Realtime • Điều phối Bếp KDS • Tự động trừ kho định lượng BOM"
    mp2.alignment = PP_ALIGN.CENTER
    mp2.font.name = "Segoe UI"
    mp2.font.size = Pt(12)
    mp2.font.color.rgb = TEXT_MUTED
    mp2.space_before = Pt(6)

    info_cards = [
        ("👤 SINH VIÊN THỰC HIỆN", [
            ("Họ và tên:", "Nguyễn Hoàng Hưng"),
            ("MSSV:", "501250384"),
            ("Khóa đào tạo:", "Khóa K25"),
            ("Chuyên ngành:", "Công nghệ Phần mềm")
        ], VIBRANT_BLUE, RGBColor(238, 242, 255)),
        ("👨‍🏫 GIẢNG VIÊN HƯỚNG DẪN", [
            ("Giảng viên:", "ThS. Nguyễn Minh Hải"),
            ("Đơn vị:", "Khoa CNTT – Điện tử"),
            ("Môn học:", "Chuyên Đề Back-End"),
            ("Thời gian:", "Tháng 10 / 2025")
        ], EMERALD, RGBColor(236, 253, 245)),
        ("⚡ CÔNG NGHỆ CHÍNH", [
            ("Môi trường:", "Node.js v20 & Express.js"),
            ("Cơ sở dữ liệu:", "MongoDB & Mongoose"),
            ("Thời gian thực:", "Socket.io WebSockets"),
            ("Xác thực:", "JWT Bearer & Bcrypt")
        ], GOLD, RGBColor(254, 243, 199))
    ]

    cx = 0.9
    for c_title, rows, c_color, c_bg in info_cards:
        c_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx), Inches(3.8), Inches(3.68), Inches(2.75))
        c_box.fill.solid()
        c_box.fill.fore_color.rgb = WHITE
        c_box.line.color.rgb = c_color
        c_box.line.width = Pt(1.5)

        ipill = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(cx + 0.2), Inches(3.98), Inches(3.28), Inches(0.38))
        ipill.fill.solid()
        ipill.fill.fore_color.rgb = c_bg
        ipill.line.fill.background()
        iptf = ipill.text_frame
        iptf.margin_top = Inches(0.06)
        ipp = iptf.paragraphs[0]
        ipp.text = c_title
        ipp.alignment = PP_ALIGN.CENTER
        ipp.font.name = "Segoe UI"
        ipp.font.size = Pt(11)
        ipp.font.bold = True
        ipp.font.color.rgb = c_color

        tx_info = slide1.shapes.add_textbox(Inches(cx + 0.25), Inches(4.5), Inches(3.18), Inches(1.9))
        itf = tx_info.text_frame
        itf.margin_top = Inches(0)
        itf.margin_left = Inches(0)
        for idx, (label, val) in enumerate(rows):
            p = itf.paragraphs[0] if idx == 0 else itf.add_paragraph()
            p.text = f"{label} "
            p.font.name = "Segoe UI"
            p.font.size = Pt(11)
            p.font.color.rgb = TEXT_MUTED
            if idx > 0:
                p.space_before = Pt(8)
            run = p.add_run()
            run.text = val
            run.font.bold = True
            run.font.color.rgb = TEXT_DARK

        cx += 3.93

    add_footer(slide1, 1)

    # ============================================================
    # SLIDE 2: ĐẶT VẤN ĐỀ & MỤC TIÊU ĐỀ TÀI (CHƯƠNG 1) - GỌN, ÍT CHỮ
    # ============================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide2, LIGHT_BG)
    add_header(slide2, "Chương 1: Giới Thiệu Tổng Quan", "Thực Trạng F&B & Mục Tiêu Hệ Thống", "Chuyển đổi quy trình thủ công sang tự động hóa thông minh với Back-End tập trung", tag_width=3.2)
    add_footer(slide2, 2)

    b_left = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(4.35))
    b_left.fill.solid()
    b_left.fill.fore_color.rgb = WHITE
    b_left.line.color.rgb = RGBColor(254, 202, 202)
    b_left.line.width = Pt(1.5)
    tf_l = b_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_top = Inches(0.2)
    tf_l.margin_left = Inches(0.25)
    tf_l.margin_right = Inches(0.25)

    p_lh = tf_l.paragraphs[0]
    p_lh.text = "⚠️ 4 BẤT CẬP QUY TRÌNH TRUYỀN THỐNG"
    p_lh.font.name = "Segoe UI"
    p_lh.font.size = Pt(12.5)
    p_lh.font.bold = True
    p_lh.font.color.rgb = RED

    pains = [
        ("Tắc nghẽn giờ cao điểm", "Ghi bill giấy chậm trễ, khách chờ lâu, giảm vòng quay bàn."),
        ("Sai sót khi truyền đạt", "Chữ viết tay khó đọc, dễ nhầm món và quên ghi chú chế biến."),
        ("Thất thoát nguyên vật liệu", "Không có công thức định lượng (BOM) gắn liền với tồn kho."),
        ("Thiếu số liệu tức thời", "Không nắm được doanh số, tồn kho và công nợ theo thời gian thực.")
    ]
    for title, desc in pains:
        p_item = tf_l.add_paragraph()
        p_item.text = f"• {title}: "
        p_item.font.name = "Segoe UI"
        p_item.font.size = Pt(10.5)
        p_item.font.bold = True
        p_item.font.color.rgb = TEXT_DARK
        p_item.space_before = Pt(10)
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    b_right = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.35), Inches(5.7), Inches(4.35))
    b_right.fill.solid()
    b_right.fill.fore_color.rgb = WHITE
    b_right.line.color.rgb = RGBColor(167, 243, 208)
    b_right.line.width = Pt(1.5)
    tf_r = b_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_top = Inches(0.2)
    tf_r.margin_left = Inches(0.25)
    tf_r.margin_right = Inches(0.25)

    p_rh = tf_r.paragraphs[0]
    p_rh.text = "🎯 4 MỤC TIÊU GIẢI PHÁP ROYAL BISTRO"
    p_rh.font.name = "Segoe UI"
    p_rh.font.size = Pt(12.5)
    p_rh.font.bold = True
    p_rh.font.color.rgb = EMERALD

    goals = [
        ("Xác thực & RBAC 3 vai trò", "Phân quyền chặt chẽ Admin, Thu ngân/Phục vụ và Đầu bếp."),
        ("Đồng bộ Realtime <50ms", "Truyền order ngay tức thì từ bàn ăn xuống màn hình Bếp KDS."),
        ("Tự động trừ kho theo BOM", "Thuật toán trừ hao phí nguyên liệu chuẩn xác khi bếp nấu."),
        ("Thanh toán QR & CRM", "Sinh mã VietQR động, tự động tích lũy điểm thưởng hội viên.")
    ]
    for title, desc in goals:
        p_item = tf_r.add_paragraph()
        p_item.text = f"• {title}: "
        p_item.font.name = "Segoe UI"
        p_item.font.size = Pt(10.5)
        p_item.font.bold = True
        p_item.font.color.rgb = TEXT_DARK
        p_item.space_before = Pt(10)
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    users = [
        ("👑 Quản trị viên (Admin)", "Toàn quyền thực đơn, kho, nhân sự, báo cáo", VIBRANT_BLUE),
        ("💳 Thu ngân / Phục vụ", "Mở bàn, tiếp nhận order POS, thanh toán", SKY_BLUE),
        ("🍳 Đầu bếp (Kitchen/Chef)", "Xem hàng đợi KDS, nấu món, kích hoạt trừ kho", GOLD),
        ("📱 Khách hàng (Customer)", "Quét mã QR tại bàn, xem menu, gọi món", EMERALD)
    ]
    u_left = 0.8
    for u_title, u_sub, u_col in users:
        u_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(u_left), Inches(5.85), Inches(2.78), Inches(1.05))
        u_box.fill.solid()
        u_box.fill.fore_color.rgb = WHITE
        u_box.line.color.rgb = CARD_BORDER
        u_box.line.width = Pt(1)
        u_tf = u_box.text_frame
        u_tf.word_wrap = True
        u_tf.margin_top = Inches(0.12)
        u_tf.margin_left = Inches(0.15)
        u_tf.margin_right = Inches(0.15)
        
        up1 = u_tf.paragraphs[0]
        up1.text = u_title
        up1.font.name = "Segoe UI"
        up1.font.size = Pt(10)
        up1.font.bold = True
        up1.font.color.rgb = u_col

        up2 = u_tf.add_paragraph()
        up2.text = u_sub
        up2.font.name = "Segoe UI"
        up2.font.size = Pt(8.5)
        up2.font.color.rgb = TEXT_MUTED
        up2.space_before = Pt(2)
        u_left += 2.97

    # ============================================================
    # SLIDE 3: CÔNG NGHỆ CỐT LÕI & KIẾN TRÚC TỔNG THỂ (CHƯƠNG 1 & 2)
    # ============================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide3, LIGHT_BG)
    add_header(slide3, "Chương 1 & 2: Công Nghệ & Kiến Trúc", "Công Nghệ Cốt Lõi & Sơ Đồ Kiến Trúc Tổng Thể", "Kiến trúc máy chủ Client-Server phân tầng, xử lý Non-blocking I/O và WebSocket hai chiều", tag_width=3.6)
    add_footer(slide3, 3)

    tech_cards = [
        ("🟢 Node.js v20 & Express v4.19", "Non-blocking I/O, Event-driven, định tuyến RESTful API cực nhanh, chịu tải cao."),
        ("🍃 MongoDB v7.0 & Mongoose v8.3", "NoSQL linh hoạt, 12 Collections chuẩn hóa, đồng nhất dữ liệu JSON từ CSDL đến Client."),
        ("⚡ Socket.io v4.7 Realtime", "Kênh truyền thông WebSocket 2 chiều, phản hồi tức thời (<50ms) giữa POS và Bếp KDS."),
        ("🔐 JWT & Bcrypt (10 Rounds)", "Xác thực phi trạng thái (Stateless), băm mật khẩu một chiều, phân quyền RBAC đa cấp.")
    ]
    t_top = 1.35
    for t_name, t_desc in tech_cards:
        t_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(t_top), Inches(5.1), Inches(1.25))
        t_box.fill.solid()
        t_box.fill.fore_color.rgb = WHITE
        t_box.line.color.rgb = CARD_BORDER
        t_box.line.width = Pt(1)
        
        ttf = t_box.text_frame
        ttf.word_wrap = True
        ttf.margin_top = Inches(0.12)
        ttf.margin_left = Inches(0.18)
        ttf.margin_right = Inches(0.18)

        tp1 = ttf.paragraphs[0]
        tp1.text = t_name
        tp1.font.name = "Segoe UI"
        tp1.font.size = Pt(11)
        tp1.font.bold = True
        tp1.font.color.rgb = PRIMARY

        tp2 = ttf.add_paragraph()
        tp2.text = t_desc
        tp2.font.name = "Segoe UI"
        tp2.font.size = Pt(9.5)
        tp2.font.color.rgb = TEXT_MUTED
        tp2.space_before = Pt(3)

        t_top += 1.42

    img_arch_path = os.path.join(media_dir, "image15.png")
    if os.path.exists(img_arch_path):
        diag_bg = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.2), Inches(1.35), Inches(6.333), Inches(5.55))
        diag_bg.fill.solid()
        diag_bg.fill.fore_color.rgb = WHITE
        diag_bg.line.color.rgb = CARD_BORDER
        diag_bg.line.width = Pt(1)
        
        cap_box = slide3.shapes.add_textbox(Inches(6.4), Inches(1.48), Inches(5.9), Inches(0.35))
        cap_tf = cap_box.text_frame
        cap_p = cap_tf.paragraphs[0]
        cap_p.text = "Sơ đồ kiến trúc tổng thể hệ thống Royal Bistro (Hình 2.1)"
        cap_p.font.name = "Segoe UI"
        cap_p.font.size = Pt(11)
        cap_p.font.bold = True
        cap_p.font.color.rgb = PRIMARY

        slide3.shapes.add_picture(img_arch_path, Inches(6.35), Inches(1.95), width=Inches(6.033))

    # ============================================================
    # SLIDE 4: TỔ CHỨC MONOREPO & 9 PHÂN HỆ NGHIỆP VỤ (CHƯƠNG 2)
    # ============================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide4, LIGHT_BG)
    add_header(slide4, "Chương 2: Thiết Kế Kiến Trúc", "Cấu Trúc Monorepo & 9 Phân Hệ Nghiệp Vụ", "Tổ chức mã nguồn khoa học, phân định rành mạch giữa Controller, Middleware, Model và Route", tag_width=3.2)
    add_footer(slide4, 4)

    left_m = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(4.5), Inches(5.55))
    left_m.fill.solid()
    left_m.fill.fore_color.rgb = WHITE
    left_m.line.color.rgb = CARD_BORDER
    left_m.line.width = Pt(1)
    tf_m = left_m.text_frame
    tf_m.word_wrap = True
    tf_m.margin_top = Inches(0.2)
    tf_m.margin_left = Inches(0.25)
    tf_m.margin_right = Inches(0.25)

    pm_h = tf_m.paragraphs[0]
    pm_h.text = "📁 PHÂN TẦNG MÃ NGUỒN (BE/src)"
    pm_h.font.name = "Segoe UI"
    pm_h.font.size = Pt(12)
    pm_h.font.bold = True
    pm_h.font.color.rgb = PRIMARY

    dirs = [
        ("config/", "Kết nối MongoDB, cấu hình port và JWT Secret."),
        ("constants/", "Định nghĩa role: admin, cashier, kitchen."),
        ("controllers/", "Xử lý nghiệp vụ chính của toàn hệ thống."),
        ("middlewares/", "Xác thực JWT, phân quyền RBAC, log và bắt lỗi."),
        ("models/", "Lược đồ Mongoose Schema & ràng buộc dữ liệu."),
        ("routes/", "Định tuyến >30 RESTful API endpoints chuẩn mực."),
        ("utils/", "Tiện ích: asyncHandler, socket emit, seeder CSDL."),
        ("server.js", "Entry point, khởi tạo Express HTTP & Socket.io.")
    ]
    for d_name, d_desc in dirs:
        pd = tf_m.add_paragraph()
        pd.text = f"• {d_name} "
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(9.5)
        pd.font.bold = True
        pd.font.color.rgb = VIBRANT_BLUE
        pd.space_before = Pt(5)
        run = pd.add_run()
        run.text = d_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    rows = 10
    cols = 3
    table_shape = slide4.shapes.add_table(rows, cols, Inches(5.6), Inches(1.35), Inches(6.933), Inches(5.55))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(1.9)
    tbl.columns[1].width = Inches(1.8)
    tbl.columns[2].width = Inches(3.233)

    headers = ["Module Nghiệp Vụ", "API Tuyến Đường", "Chức Năng Trọng Tâm"]
    for c_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        cp = cell.text_frame.paragraphs[0]
        cp.text = h_text
        cp.font.name = "Segoe UI"
        cp.font.size = Pt(10)
        cp.font.bold = True
        cp.font.color.rgb = WHITE
        cell.text_frame.margin_top = Inches(0.06)
        cell.text_frame.margin_bottom = Inches(0.06)

    mod_data = [
        ("Xác thực & Nhân sự", "/api/v1/auth", "Đăng nhập, đăng ký, JWT token, 1-Click Demo Login"),
        ("Quản lý Bàn ăn", "/api/v1/tables", "Sơ đồ mặt bằng, sức chứa, trạng thái động"),
        ("Thực đơn & Món ăn", "/api/v1/dishes", "Phân loại danh mục, đơn giá, chuyển trạng thái bán"),
        ("Bán hàng POS", "/api/v1/orders", "Gọi món tại bàn, ghi chú chế biến, mức ưu tiên"),
        ("Điều phối Bếp KDS", "/api/v1/orders/:id/status", "Tiếp nhận món, cập nhật nấu xong, trừ kho BOM"),
        ("Quản lý Đặt bàn", "/api/v1/reservations", "Đặt trước bàn, kiểm tra bàn trống, nhận cọc"),
        ("Kho & Định lượng BOM", "/api/v1/inventory", "Theo dõi tồn kho, cảnh báo cạn kiệt, định mức món"),
        ("Khách hàng & CRM", "/api/v1/customers", "Tích lũy điểm thưởng theo giá trị hóa đơn"),
        ("Báo cáo & Thống kê", "/api/v1/reports", "Doanh thu ngày/tháng, Top 5 món bán chạy nhất")
    ]
    for r_idx, (m_name, m_route, m_func) in enumerate(mod_data, start=1):
        bg_col = WHITE if r_idx % 2 != 0 else CARD_ALT
        for c_idx, text_val in enumerate([m_name, m_route, m_func]):
            cell = tbl.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_col
            cp = cell.text_frame.paragraphs[0]
            cp.text = text_val
            cp.font.name = "Segoe UI"
            cp.font.size = Pt(9.5)
            if c_idx == 0:
                cp.font.bold = True
                cp.font.color.rgb = TEXT_DARK
            elif c_idx == 1:
                cp.font.color.rgb = VIBRANT_BLUE
                cp.font.name = "Consolas"
                cp.font.size = Pt(8.5)
            else:
                cp.font.color.rgb = TEXT_MUTED
            cell.text_frame.margin_top = Inches(0.04)
            cell.text_frame.margin_bottom = Inches(0.04)

    # ============================================================
    # SLIDE 5: THIẾT KẾ CƠ SỞ DỮ LIỆU - 12 COLLECTIONS (CHƯƠNG 2)
    # ============================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide5, LIGHT_BG)
    add_header(slide5, "Chương 2: Cơ Sở Dữ Liệu", "Thiết Kế Cơ Sở Dữ Liệu NoSQL 12 Collections", "Tối ưu hóa cấu trúc dữ liệu NoSQL với Mongoose ODM, giải quyết triệt để quan hệ định mức n-n (BOM)", tag_width=3.0)
    add_footer(slide5, 5)

    left_c = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.2), Inches(5.55))
    left_c.fill.solid()
    left_c.fill.fore_color.rgb = WHITE
    left_c.line.color.rgb = CARD_BORDER
    left_c.line.width = Pt(1)
    tf_c = left_c.text_frame
    tf_c.word_wrap = True
    tf_c.margin_top = Inches(0.2)
    tf_c.margin_left = Inches(0.25)
    tf_c.margin_right = Inches(0.25)

    pc_h = tf_c.paragraphs[0]
    pc_h.text = "🗄️ 4 NHÓM THỰC THỂ DỮ LIỆU CHÍNH"
    pc_h.font.name = "Segoe UI"
    pc_h.font.size = Pt(12)
    pc_h.font.bold = True
    pc_h.font.color.rgb = PRIMARY

    c_groups = [
        ("1. Người Dùng & Bàn Ăn", "NguoiDung (Admin, Cashier, Kitchen), BanAn & KhuVuc (Sơ đồ mặt bằng, trạng thái bàn động)."),
        ("2. Thực Đơn & Gọi Món", "MonAn, LoaiMon (Danh mục & giá bán), DatMon (Chi tiết gọi món POS, ghi chú đặc biệt)."),
        ("3. Kho & Công Thức BOM", "NguyenLieu (Tồn kho thực tế), MonAnNguyenLieu (Bảng định lượng liên kết n-n trừ kho tự động)."),
        ("4. Đặt Bàn, Hóa Đơn & CRM", "DatBan (Khách hẹn trước), HoaDon (Thanh toán tiền mặt/VietQR), KhachHang (Tích điểm loyalty).")
    ]
    for grp_t, grp_d in c_groups:
        pg = tf_c.add_paragraph()
        pg.text = f"• {grp_t}: "
        pg.font.name = "Segoe UI"
        pg.font.size = Pt(10.5)
        pg.font.bold = True
        pg.font.color.rgb = VIBRANT_BLUE
        pg.space_before = Pt(10)
        run = pg.add_run()
        run.text = grp_d
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    img_erd_path = os.path.join(media_dir, "image19.png")
    if os.path.exists(img_erd_path):
        erd_bg = slide5.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.3), Inches(1.35), Inches(6.233), Inches(5.55))
        erd_bg.fill.solid()
        erd_bg.fill.fore_color.rgb = WHITE
        erd_bg.line.color.rgb = CARD_BORDER
        erd_bg.line.width = Pt(1)

        cap_erd = slide5.shapes.add_textbox(Inches(6.5), Inches(1.48), Inches(5.8), Inches(0.35))
        cap_etf = cap_erd.text_frame
        cap_ep = cap_etf.paragraphs[0]
        cap_ep.text = "Sơ đồ thực thể quan hệ (ERD) 12 Collections MongoDB (Hình 2.5)"
        cap_ep.font.name = "Segoe UI"
        cap_ep.font.size = Pt(11)
        cap_ep.font.bold = True
        cap_ep.font.color.rgb = PRIMARY

        slide5.shapes.add_picture(img_erd_path, Inches(6.45), Inches(1.95), width=Inches(5.933))

    # ============================================================
    # SLIDE 6: CÁC LUỒNG NGHIỆP VỤ TRỌNG TÂM (CHƯƠNG 2) - GỌN + LƯU ĐỒ
    # ============================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide6, LIGHT_BG)
    add_header(slide6, "Chương 2: Luồng Xử Lý Nghiệp Vụ", "3 Luồng Nghiệp Vụ Trọng Tâm Của Hệ Thống", "Quy trình xác thực RBAC, gọi món POS - KDS Realtime và thuật toán trừ kho BOM tự động", tag_width=3.6)
    add_footer(slide6, 6)

    flows_data = [
        ("🔐 LUỒNG 1: XÁC THỰC & PHÂN QUYỀN RBAC", [
            ("Đăng nhập:", "POST /api/v1/auth/login kèm email/mật khẩu."),
            ("So khớp mật khẩu:", "Bcrypt kiểm tra băm 10 salt rounds."),
            ("Cấp JWT Token:", "Thời hạn 24h chứa vai trò (Admin/Cashier/Kitchen)."),
            ("Chặn truy cập:", "Middleware RBAC bảo vệ từng API endpoint.")
        ], VIBRANT_BLUE),
        ("⚡ LUỒNG 2: GỌI MÓN POS -> BẾP KDS REALTIME", [
            ("Lên đơn tại bàn:", "Nhân viên chọn món, ghi chú cay/ngọt, gửi đơn."),
            ("Đổi trạng thái bàn:", "Chuyển bàn sang 'co_khach' tức thời."),
            ("Phát sóng Socket.io:", "Gửi sự kiện 'order:new' xuống màn hình bếp."),
            ("Bếp nhận <50ms:", "Chuông reo báo hiệu, xếp vào hàng đợi ưu tiên.")
        ], GOLD),
        ("🍳 LUỒNG 3: ĐIỀU PHỐI BẾP & TRỪ KHO BOM", [
            ("Đầu bếp bấm 'Nấu':", "Gửi PUT /api/v1/orders/:id/status ('dang_nau')."),
            ("Truy vấn bảng BOM:", "Lấy công thức tiêu hao từ MonAnNguyenLieu."),
            ("Tính toán hao phí:", "soLuongCanTru = bom.dinh_muc * don.so_luong."),
            ("Trừ kho an toàn:", "Trừ tồn kho và cảnh báo khi chạm ngưỡng tối thiểu.")
        ], EMERALD)
    ]

    fl_top = 1.35
    for fl_title, steps, fl_col in flows_data:
        fl_box = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(fl_top), Inches(5.8), Inches(1.75))
        fl_box.fill.solid()
        fl_box.fill.fore_color.rgb = WHITE
        fl_box.line.color.rgb = fl_col
        fl_box.line.width = Pt(1.5)
        
        fl_tf = fl_box.text_frame
        fl_tf.word_wrap = True
        fl_tf.margin_top = Inches(0.1)
        fl_tf.margin_left = Inches(0.18)
        fl_tf.margin_right = Inches(0.18)

        fl_p = fl_tf.paragraphs[0]
        fl_p.text = fl_title
        fl_p.font.name = "Segoe UI"
        fl_p.font.size = Pt(10.5)
        fl_p.font.bold = True
        fl_p.font.color.rgb = fl_col

        for s_title, s_desc in steps:
            fl_ps = fl_tf.add_paragraph()
            fl_ps.text = f"• {s_title} "
            fl_ps.font.name = "Segoe UI"
            fl_ps.font.size = Pt(9)
            fl_ps.font.bold = True
            fl_ps.font.color.rgb = TEXT_DARK
            fl_ps.space_before = Pt(2)
            run = fl_ps.add_run()
            run.text = s_desc
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

        fl_top += 1.88

    img_flow_path = os.path.join(media_dir, "image21.png")
    if os.path.exists(img_flow_path):
        flow_bg = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.9), Inches(1.35), Inches(5.633), Inches(5.55))
        flow_bg.fill.solid()
        flow_bg.fill.fore_color.rgb = WHITE
        flow_bg.line.color.rgb = CARD_BORDER
        flow_bg.line.width = Pt(1)

        cap_fl = slide6.shapes.add_textbox(Inches(7.1), Inches(1.48), Inches(5.2), Inches(0.35))
        cap_ftf = cap_fl.text_frame
        cap_fp = cap_ftf.paragraphs[0]
        cap_fp.text = "Lưu đồ gọi món POS & điều phối Bếp KDS Real-time (Hình 2.7)"
        cap_fp.font.name = "Segoe UI"
        cap_fp.font.size = Pt(10.5)
        cap_fp.font.bold = True
        cap_fp.font.color.rgb = PRIMARY

        slide6.shapes.add_picture(img_flow_path, Inches(7.05), Inches(1.95), width=Inches(5.333))

    # ============================================================
    # SLIDE 7: GIAO DIỆN & TÍNH NĂNG THỰC NGHIỆM (CHƯƠNG 3) - ẢNH NỔI BẬT
    # ============================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide7, LIGHT_BG)
    add_header(slide7, "Chương 3: Cài Đặt Thực Nghiệm", "Giao Diện & Tính Năng Thực Nghiệm Trực Quan", "Kiểm thử thực tế hệ thống trên các màn hình Dashboard quản trị, sơ đồ bàn ăn và điều phối Bếp KDS", tag_width=3.4)
    add_footer(slide7, 7)

    ui_cards = [
        ("BẢNG ĐIỀU KHIỂN QUẢN TRỊ (DASHBOARD)", "image24.png", "• Thống kê doanh số theo ngày/ca làm việc.\n• Top 5 món bán chạy nhất & tỷ lệ lấp đầy bàn."),
        ("SƠ ĐỒ BÀN ĂN TRẠNG THÁI ĐỘNG", "image25.png", "• Hiển thị trực quan mặt bằng theo khu vực.\n• Đổi màu động (trống, có khách, đã đặt), chuyển bàn."),
        ("MÀN HÌNH BẾP KDS REAL-TIME", "image27.png", "• Nhận đơn tức thì (<50ms), chuông reo báo hiệu.\n• Cập nhật tiến độ 1 chạm & kích hoạt trừ kho BOM.")
    ]

    card_x = 0.8
    for ui_title, img_name, ui_desc in ui_cards:
        u_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(card_x), Inches(1.35), Inches(3.75), Inches(5.55))
        u_box.fill.solid()
        u_box.fill.fore_color.rgb = WHITE
        u_box.line.color.rgb = CARD_BORDER
        u_box.line.width = Pt(1)
        
        u_tf = u_box.text_frame
        u_tf.word_wrap = True
        u_tf.margin_top = Inches(0.15)
        u_tf.margin_left = Inches(0.18)
        u_tf.margin_right = Inches(0.18)

        up = u_tf.paragraphs[0]
        up.text = ui_title
        up.font.name = "Segoe UI"
        up.font.size = Pt(10)
        up.font.bold = True
        up.font.color.rgb = PRIMARY

        img_p = os.path.join(media_dir, img_name)
        if os.path.exists(img_p):
            slide7.shapes.add_picture(img_p, Inches(card_x + 0.15), Inches(1.9), width=Inches(3.45))

        desc_box = slide7.shapes.add_textbox(Inches(card_x + 0.15), Inches(4.3), Inches(3.45), Inches(2.4))
        dtf = desc_box.text_frame
        dtf.word_wrap = True
        dtf.margin_left = Inches(0)
        dtf.margin_top = Inches(0)
        dp = dtf.paragraphs[0]
        dp.text = ui_desc
        dp.font.name = "Segoe UI"
        dp.font.size = Pt(9.5)
        dp.font.color.rgb = TEXT_MUTED

        card_x += 3.99

    # ============================================================
    # SLIDE 8: XỬ LÝ LỖI & GIÁM SÁT HỆ THỐNG (CHƯƠNG 4) - NGẮN GỌN
    # ============================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide8, LIGHT_BG)
    add_header(slide8, "Chương 4: Xử Lý Lỗi & Giám Sát", "Kiến Trúc Xử Lý Lỗi Tập Trung & Giám Sát Hiệu Năng", "Cơ chế bắt lỗi bất đồng bộ asyncHandler, bảng mã lỗi HTTP chuẩn hóa và Request Logger đo độ trễ", tag_width=3.6)
    add_footer(slide8, 8)

    left_e = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.5), Inches(5.55))
    left_e.fill.solid()
    left_e.fill.fore_color.rgb = WHITE
    left_e.line.color.rgb = CARD_BORDER
    left_e.line.width = Pt(1)
    tf_e = left_e.text_frame
    tf_e.word_wrap = True
    tf_e.margin_top = Inches(0.2)
    tf_e.margin_left = Inches(0.25)
    tf_e.margin_right = Inches(0.25)

    pe_h = tf_e.paragraphs[0]
    pe_h.text = "🛡️ 4 CƠ CHẾ XỬ LÝ LỖI CHUYÊN SÂU"
    pe_h.font.name = "Segoe UI"
    pe_h.font.size = Pt(12)
    pe_h.font.bold = True
    pe_h.font.color.rgb = PRIMARY

    err_points = [
        ("Tiện ích asyncHandler", "Bọc quanh các controller async/await, loại bỏ mã lặp try/catch thừa, tự động đẩy lỗi về middleware."),
        ("Custom Error kế thừa AppError", "Định nghĩa mã lỗi chuẩn HTTP: LoiYeuCauKhongHopLe, LoiChuaXacThuc, LoiKhongTimThay..."),
        ("Middleware xử lý lỗi tập trung", "Không bao giờ để lộ stack trace nhạy cảm khi chạy production; trả về JSON đồng nhất."),
        ("Middleware Request Logger", "Đo thời gian phản hồi (ms) cho từng request HTTP, phát hiện sớm các API bị nghẽn cổ chai.")
    ]
    for ep_name, ep_desc in err_points:
        pe = tf_e.add_paragraph()
        pe.text = f"• {ep_name}: "
        pe.font.name = "Segoe UI"
        pe.font.size = Pt(10)
        pe.font.bold = True
        pe.font.color.rgb = VIBRANT_BLUE
        pe.space_before = Pt(10)
        run = pe.add_run()
        run.text = ep_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    t_shape8 = slide8.shapes.add_table(7, 3, Inches(6.6), Inches(1.35), Inches(5.933), Inches(5.55))
    tbl8 = t_shape8.table
    tbl8.columns[0].width = Inches(1.4)
    tbl8.columns[1].width = Inches(1.9)
    tbl8.columns[2].width = Inches(2.633)

    h8 = ["HTTP Code", "Custom Error Class", "Tình Huống Xảy Ra"]
    for c_idx, h_text in enumerate(h8):
        cell = tbl8.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        cp = cell.text_frame.paragraphs[0]
        cp.text = h_text
        cp.font.name = "Segoe UI"
        cp.font.size = Pt(10)
        cp.font.bold = True
        cp.font.color.rgb = WHITE
        cell.text_frame.margin_top = Inches(0.06)

    codes_data = [
        ("400 Bad Request", "LoiYeuCauKhongHopLe", "Thiếu tham số bắt buộc, số lượng món <= 0"),
        ("401 Unauthorized", "LoiChuaXacThuc", "Chưa gửi token, token hết hạn, sai mật khẩu"),
        ("403 Forbidden", "LoiKhongCoQuyen", "Nhân viên cố truy cập chức năng của Admin"),
        ("404 Not Found", "LoiKhongTimThay", "Mã bàn, món ăn hoặc đơn hàng không tồn tại"),
        ("409 Conflict", "LoiXungDotDuLieu", "Email nhân viên đã tồn tại; bàn đang có khách"),
        ("500 Server Error", "LoiHeThong", "Lỗi máy chủ nội bộ hoặc lỗi kết nối MongoDB")
    ]
    for r_idx, (code, err_cls, cond) in enumerate(codes_data, start=1):
        bg = WHITE if r_idx % 2 != 0 else CARD_ALT
        for c_idx, val in enumerate([code, err_cls, cond]):
            cell = tbl8.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cp = cell.text_frame.paragraphs[0]
            cp.text = val
            cp.font.name = "Segoe UI"
            cp.font.size = Pt(9.5)
            if c_idx == 0:
                cp.font.bold = True
                cp.font.color.rgb = RED
            elif c_idx == 1:
                cp.font.name = "Consolas"
                cp.font.size = Pt(8.5)
                cp.font.color.rgb = VIBRANT_BLUE
            else:
                cp.font.color.rgb = TEXT_MUTED
            cell.text_frame.margin_top = Inches(0.05)

    # ============================================================
    # SLIDE 9: BẢO MẬT & KIỂM THỬ CHỊU TẢI (CHƯƠNG 4) - NGẮN GỌN
    # ============================================================
    slide9 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide9, LIGHT_BG)
    add_header(slide9, "Chương 4: Bảo Mật & Kiểm Thử", "Bảo Mật Đa Tầng & Kết Quả Kiểm Thử Chịu Tải API", "Triển khai Helmet, CORS, Rate Limiting, Bcrypt, JWT và kiểm thử chịu tải với 100% tỷ lệ phản hồi thành công", tag_width=3.4)
    add_footer(slide9, 9)

    left_s = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.0), Inches(5.55))
    left_s.fill.solid()
    left_s.fill.fore_color.rgb = WHITE
    left_s.line.color.rgb = CARD_BORDER
    left_s.line.width = Pt(1)
    tf_s = left_s.text_frame
    tf_s.word_wrap = True
    tf_s.margin_top = Inches(0.2)
    tf_s.margin_left = Inches(0.25)
    tf_s.margin_right = Inches(0.25)

    ps_h = tf_s.paragraphs[0]
    ps_h.text = "🔒 5 LỚP BẢO MẬT HỆ THỐNG"
    ps_h.font.name = "Segoe UI"
    ps_h.font.size = Pt(12)
    ps_h.font.bold = True
    ps_h.font.color.rgb = PRIMARY

    sec_items = [
        ("Helmet.js Headers", "Chống các lỗ hổng web phổ biến (XSS, Clickjacking, MIME sniffing)."),
        ("CORS Whitelist", "Chỉ chấp nhận request từ domain Front-End hợp lệ."),
        ("Express Rate Limit", "100 req/15 phút (API chung) và 20 req/15 phút (Auth), chống DoS và Brute-force."),
        ("Bcrypt (10 Rounds)", "Mật khẩu băm một chiều với muối ngẫu nhiên trước khi lưu DB."),
        ("Stateless JWT 24h", "Phiên làm việc phi trạng thái, không tốn session RAM máy chủ.")
    ]
    for s_name, s_desc in sec_items:
        ps = tf_s.add_paragraph()
        ps.text = f"• {s_name}: "
        ps.font.name = "Segoe UI"
        ps.font.size = Pt(10)
        ps.font.bold = True
        ps.font.color.rgb = VIBRANT_BLUE
        ps.space_before = Pt(8)
        run = ps.add_run()
        run.text = s_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    t_shape9 = slide9.shapes.add_table(6, 5, Inches(6.1), Inches(1.35), Inches(6.433), Inches(3.6))
    tbl9 = t_shape9.table
    tbl9.columns[0].width = Inches(2.2)
    tbl9.columns[1].width = Inches(1.1)
    tbl9.columns[2].width = Inches(1.1)
    tbl9.columns[3].width = Inches(1.1)
    tbl9.columns[4].width = Inches(0.933)

    h9 = ["Kịch Bản Kiểm Thử", "Yêu Cầu", "Đồng Thời", "TG Phản Hồi", "Tỷ Lệ Lỗi"]
    for c_idx, h_text in enumerate(h9):
        cell = tbl9.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        cp = cell.text_frame.paragraphs[0]
        cp.text = h_text
        cp.font.name = "Segoe UI"
        cp.font.size = Pt(9.5)
        cp.font.bold = True
        cp.font.color.rgb = WHITE
        cell.text_frame.margin_top = Inches(0.05)

    perf_data = [
        ("Lấy DS bàn ăn (GET /tables)", "1,000 req", "50 kết nối", "18.5 ms", "0.0%"),
        ("Lấy thực đơn (GET /dishes)", "1,000 req", "50 kết nối", "22.1 ms", "0.0%"),
        ("Tạo đơn mới (POST /orders)", "500 req", "25 kết nối", "42.8 ms", "0.0%"),
        ("Cập nhật & trừ kho BOM", "500 req", "25 kết nối", "38.4 ms", "0.0%"),
        ("Thanh toán (POST /payment)", "200 req", "10 kết nối", "48.2 ms", "0.0%"),
    ]
    for r_idx, row in enumerate(perf_data, start=1):
        bg = WHITE if r_idx % 2 != 0 else CARD_ALT
        for c_idx, val in enumerate(row):
            cell = tbl9.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg
            cp = cell.text_frame.paragraphs[0]
            cp.text = val
            cp.font.name = "Segoe UI"
            cp.font.size = Pt(9.5)
            if c_idx == 0:
                cp.font.bold = True
                cp.font.color.rgb = TEXT_DARK
            elif c_idx == 3 or c_idx == 4:
                cp.font.bold = True
                cp.font.color.rgb = EMERALD
            else:
                cp.font.color.rgb = TEXT_MUTED
            cell.text_frame.margin_top = Inches(0.04)

    eval_box = slide9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.1), Inches(5.15), Inches(6.433), Inches(1.75))
    eval_box.fill.solid()
    eval_box.fill.fore_color.rgb = RGBColor(236, 253, 245)
    eval_box.line.color.rgb = RGBColor(167, 243, 208)
    eval_box.line.width = Pt(1.5)
    
    etf = eval_box.text_frame
    etf.word_wrap = True
    etf.margin_top = Inches(0.12)
    etf.margin_left = Inches(0.2)
    etf.margin_right = Inches(0.2)

    ep1 = etf.paragraphs[0]
    ep1.text = "✅ ĐÁNH GIÁ CHỨC NĂNG NGHIỆP VỤ (BẢNG 3.3)"
    ep1.font.name = "Segoe UI"
    ep1.font.size = Pt(11)
    ep1.font.bold = True
    ep1.font.color.rgb = EMERALD

    ep2 = etf.add_paragraph()
    ep2.text = "• 9/9 nhóm chức năng chính được kiểm thử thực tế và đánh giá ĐẠT CHUẨN 100%."
    ep2.font.name = "Segoe UI"
    ep2.font.size = Pt(9.5)
    ep2.font.color.rgb = TEXT_DARK
    ep2.space_before = Pt(3)

    ep3 = etf.add_paragraph()
    ep3.text = "• Thời gian phản hồi API trung bình dưới 50ms, độ tin cậy tuyệt đối không phát sinh lỗi ngoại lệ."
    ep3.font.name = "Segoe UI"
    ep3.font.size = Pt(9.5)
    ep3.font.color.rgb = TEXT_DARK
    ep3.space_before = Pt(2)

    # ============================================================
    # SLIDE 10: TỔNG KẾT & HƯỚNG PHÁT TRIỂN (CHƯƠNG 5) - TRANG RIÊNG
    # ============================================================
    slide10 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide10, LIGHT_BG)
    add_header(slide10, "Chương 5: Kết Luận & Hướng Phát Triển", "Tổng Kết Đề Tài & Định Hướng Nâng Cấp", "Đánh giá các kết quả đạt được và lộ trình phát triển hệ thống trong tương lai", tag_width=4.0)
    add_footer(slide10, 10)

    # Left Card: Achievements (Full Height, Rộng Rãi)
    left_10 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.55))
    left_10.fill.solid()
    left_10.fill.fore_color.rgb = WHITE
    left_10.line.color.rgb = RGBColor(167, 243, 208)
    left_10.line.width = Pt(1.5)
    tf_10l = left_10.text_frame
    tf_10l.word_wrap = True
    tf_10l.margin_top = Inches(0.25)
    tf_10l.margin_left = Inches(0.28)
    tf_10l.margin_right = Inches(0.28)

    p10_lh = tf_10l.paragraphs[0]
    p10_lh.text = "🌟 4 THÀNH TỰU CỐT LÕI ĐÃ ĐẠT ĐƯỢC"
    p10_lh.font.name = "Segoe UI"
    p10_lh.font.size = Pt(13)
    p10_lh.font.bold = True
    p10_lh.font.color.rgb = EMERALD

    achieves = [
        ("Kiến trúc Back-End chuẩn mực", "Xây dựng máy chủ Node.js & Express phân tầng rành mạch, kiểm soát hoàn toàn hơn 30 RESTful API endpoints."),
        ("CSDL 12 Collections hoàn thiện", "Mô hình hóa Mongoose ODM quan hệ chặt chẽ, giải quyết triệt để bài toán công thức định mức BOM n-n."),
        ("Thời gian thực Realtime vượt trội", "Ứng dụng Socket.io WebSockets phát sóng tức thời giữa POS và Bếp KDS với độ trễ <50ms."),
        ("Bảo mật & Xử lý lỗi vững chắc", "asyncHandler loại bỏ try/catch thừa, kiểm soát lỗi tập trung, tích hợp Helmet, CORS, Rate Limit và JWT.")
    ]
    for ach_t, ach_d in achieves:
        pach = tf_10l.add_paragraph()
        pach.text = f"• {ach_t}: "
        pach.font.name = "Segoe UI"
        pach.font.size = Pt(10.5)
        pach.font.bold = True
        pach.font.color.rgb = TEXT_DARK
        pach.space_before = Pt(14)
        run = pach.add_run()
        run.text = ach_d
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Right Card: Future Roadmap (Full Height, Rộng Rãi)
    right_10 = slide10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.35), Inches(5.7), Inches(5.55))
    right_10.fill.solid()
    right_10.fill.fore_color.rgb = WHITE
    right_10.line.color.rgb = RGBColor(253, 230, 138)
    right_10.line.width = Pt(1.5)
    tf_10r = right_10.text_frame
    tf_10r.word_wrap = True
    tf_10r.margin_top = Inches(0.25)
    tf_10r.margin_left = Inches(0.28)
    tf_10r.margin_right = Inches(0.28)

    p10_rh = tf_10r.paragraphs[0]
    p10_rh.text = "🚀 LỘ TRÌNH PHÁT TRIỂN TIẾP THEO"
    p10_rh.font.name = "Segoe UI"
    p10_rh.font.size = Pt(13)
    p10_rh.font.bold = True
    p10_rh.font.color.rgb = GOLD

    futures = [
        ("Tích hợp Redis Cache", "Lưu bộ nhớ đệm thực đơn & sơ đồ bàn ăn để giảm tải truy vấn MongoDB và tăng tốc phản hồi."),
        ("Webhook Ngân hàng tự động", "Kết nối SePAY hoặc VietQR Webhook để tự động xác nhận hóa đơn khi tiền vào tài khoản."),
        ("Kiến trúc Microservices", "Tách độc lập Service Báo cáo (Reporting) và Service Kho nguyên liệu (Inventory) để mở rộng quy mô."),
        ("Ứng dụng Trí tuệ nhân tạo (AI)", "Xây dựng mô hình AI dự báo lượng khách và định lượng tiêu hao nguyên liệu theo ngày trong tuần.")
    ]
    for fut_t, fut_d in futures:
        pfut = tf_10r.add_paragraph()
        pfut.text = f"• {fut_t}: "
        pfut.font.name = "Segoe UI"
        pfut.font.size = Pt(10.5)
        pfut.font.bold = True
        pfut.font.color.rgb = TEXT_DARK
        pfut.space_before = Pt(14)
        run = pfut.add_run()
        run.text = fut_d
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # ============================================================
    # SLIDE 11: SLIDE CẢM ƠN & HỎI ĐÁP (DEDICATED THANK YOU & Q&A)
    # ============================================================
    slide11 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide11, LIGHT_BG)

    # Top accent bar
    top_bar11 = slide11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar11.fill.solid()
    top_bar11.fill.fore_color.rgb = VIBRANT_BLUE
    top_bar11.line.fill.background()

    # Main Center Gratitude Card
    thanks_card = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.0), Inches(10.333), Inches(4.3))
    thanks_card.fill.solid()
    thanks_card.fill.fore_color.rgb = WHITE
    thanks_card.line.color.rgb = RGBColor(199, 210, 254)
    thanks_card.line.width = Pt(1.5)

    tt_frame = thanks_card.text_frame
    tt_frame.word_wrap = True
    tt_frame.margin_top = Inches(0.4)
    tt_frame.margin_left = Inches(0.5)
    tt_frame.margin_right = Inches(0.5)

    p_tbadge = tt_frame.paragraphs[0]
    p_tbadge.text = "🏛️ KHOA CÔNG NGHỆ THÔNG TIN – ĐIỆN TỬ • ITC"
    p_tbadge.alignment = PP_ALIGN.CENTER
    p_tbadge.font.name = "Segoe UI"
    p_tbadge.font.size = Pt(11)
    p_tbadge.font.bold = True
    p_tbadge.font.color.rgb = VIBRANT_BLUE

    p_thx = tt_frame.add_paragraph()
    p_thx.text = "XIN CHÂN THÀNH CẢM ƠN!"
    p_thx.alignment = PP_ALIGN.CENTER
    p_thx.font.name = "Segoe UI"
    p_thx.font.size = Pt(30)
    p_thx.font.bold = True
    p_thx.font.color.rgb = PRIMARY
    p_thx.space_before = Pt(8)

    p_sub = tt_frame.add_paragraph()
    p_sub.text = "Kính chúc Quý Thầy Cô trong Hội đồng thật nhiều sức khỏe, hạnh phúc và thành công!"
    p_sub.alignment = PP_ALIGN.CENTER
    p_sub.font.name = "Segoe UI"
    p_sub.font.size = Pt(13)
    p_sub.font.color.rgb = TEXT_MUTED
    p_sub.space_before = Pt(8)

    # Q&A Highlight inside the card
    p_qa = tt_frame.add_paragraph()
    p_qa.text = "❓ PHẦN HỎI & ĐÁP (Q & A)"
    p_qa.alignment = PP_ALIGN.CENTER
    p_qa.font.name = "Segoe UI"
    p_qa.font.size = Pt(16)
    p_qa.font.bold = True
    p_qa.font.color.rgb = GOLD
    p_qa.space_before = Pt(20)

    p_qa_desc = tt_frame.add_paragraph()
    p_qa_desc.text = "Em rất mong nhận được những câu hỏi phản biện, lời nhận xét và góp ý quý báu của Quý Thầy Cô để hoàn thiện đề tài."
    p_qa_desc.alignment = PP_ALIGN.CENTER
    p_qa_desc.font.name = "Segoe UI"
    p_qa_desc.font.size = Pt(11.5)
    p_qa_desc.font.color.rgb = TEXT_DARK
    p_qa_desc.space_before = Pt(6)

    # Bottom Contact / Author Info Card
    contact_card = slide11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(5.55), Inches(10.333), Inches(1.3))
    contact_card.fill.solid()
    contact_card.fill.fore_color.rgb = RGBColor(238, 242, 255)
    contact_card.line.color.rgb = RGBColor(199, 210, 254)
    contact_card.line.width = Pt(1)

    ctf = contact_card.text_frame
    ctf.word_wrap = True
    ctf.margin_top = Inches(0.18)
    ctf.margin_left = Inches(0.4)
    ctf.margin_right = Inches(0.4)

    cp1 = ctf.paragraphs[0]
    cp1.text = "👤 Sinh viên thực hiện: Nguyễn Hoàng Hưng  •  MSSV: 501250384  •  Lớp: K25  •  Chuyên ngành: Công nghệ Phần mềm"
    cp1.alignment = PP_ALIGN.CENTER
    cp1.font.name = "Segoe UI"
    cp1.font.size = Pt(11)
    cp1.font.bold = True
    cp1.font.color.rgb = PRIMARY

    cp2 = ctf.add_paragraph()
    cp2.text = "👨‍🏫 Giảng viên hướng dẫn: ThS. Nguyễn Minh Hải  |  Đề tài: Hệ Thống Back-End Quản Lý Nhà Hàng & Gọi Món (Royal Bistro)"
    cp2.alignment = PP_ALIGN.CENTER
    cp2.font.name = "Segoe UI"
    cp2.font.size = Pt(10.5)
    cp2.font.color.rgb = TEXT_MUTED
    cp2.space_before = Pt(4)

    add_footer(slide11, 11)

    # --- SAVE FILE ---
    output_path = r'c:\Users\LENOVO\OneDrive\Documents\My Data Sources\PhanMemQuanLyMonAn\BaoCao_TongQuan_DoAn_BackEnd_RoyalBistro.pptx'
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation with dedicated Thank You slide at: {output_path}")

if __name__ == "__main__":
    create_presentation()
