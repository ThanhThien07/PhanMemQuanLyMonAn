import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- PALETTE MÀU HIỆN ĐẠI, TƯƠNG PHẢN CAO, CHUẨN ROYAL BISTRO ---
    LIGHT_BG = RGBColor(248, 250, 252)    # Slate 50
    PRIMARY = RGBColor(30, 58, 138)       # Royal Blue 900
    VIBRANT_BLUE = RGBColor(37, 99, 235)  # Blue 600
    SKY_BLUE = RGBColor(2, 132, 199)      # Sky 600
    GOLD = RGBColor(217, 119, 6)          # Amber 600
    EMERALD = RGBColor(5, 150, 105)       # Emerald 600
    RED = RGBColor(220, 38, 38)           # Red 600
    TEXT_DARK = RGBColor(15, 23, 42)      # Slate 900
    TEXT_MUTED = RGBColor(71, 85, 105)    # Slate 600
    CARD_BORDER = RGBColor(226, 232, 240) # Slate 200
    CARD_ALT = RGBColor(241, 245, 249)    # Slate 100
    WHITE = RGBColor(255, 255, 255)

    base_dir = os.path.dirname(os.path.abspath(__file__))
    docs_dir = os.path.dirname(base_dir)
    media_dir = os.path.join(docs_dir, 'assets', 'docx_media')
    clean_screenshots_dir = os.path.join(docs_dir, 'assets', 'screenshots_clean')
    if not os.path.exists(clean_screenshots_dir):
        clean_screenshots_dir = os.path.join(docs_dir, 'assets', 'screenshots')

    TOTAL_SLIDES = 8

    def set_slide_bg(slide, color=LIGHT_BG):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, tag_text, title_text, subtitle_text="", tag_width=4.4):
        header_box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.2))
        header_box.fill.solid()
        header_box.fill.fore_color.rgb = WHITE
        header_box.line.color.rgb = CARD_BORDER
        header_box.line.width = Pt(1)

        accent_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.06))
        accent_bar.fill.solid()
        accent_bar.fill.fore_color.rgb = VIBRANT_BLUE
        accent_bar.line.fill.background()

        tag_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.1), Inches(tag_width), Inches(0.29))
        tag_shape.adjustments[0] = 0.3
        tag_shape.fill.solid()
        tag_shape.fill.fore_color.rgb = RGBColor(238, 242, 255)
        tag_shape.line.color.rgb = RGBColor(199, 210, 254)
        tag_shape.line.width = Pt(0.75)
        tf_tag = tag_shape.text_frame
        tf_tag.word_wrap = False
        tf_tag.margin_top = Inches(0.02)
        tf_tag.margin_bottom = Inches(0.02)
        tf_tag.margin_left = Inches(0.1)
        tf_tag.margin_right = Inches(0.1)
        p_tag = tf_tag.paragraphs[0]
        p_tag.text = tag_text.upper()
        p_tag.alignment = PP_ALIGN.CENTER
        p_tag.font.name = "Segoe UI"
        p_tag.font.size = Pt(10.5)
        p_tag.font.bold = True
        p_tag.font.color.rgb = VIBRANT_BLUE

        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.40), Inches(11.733), Inches(0.76))
        tf = txBox.text_frame
        tf.word_wrap = True
        tf.margin_top = Inches(0)
        tf.margin_left = Inches(0)
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Segoe UI"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = PRIMARY

        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(13)
            p2.font.color.rgb = TEXT_MUTED
            p2.space_before = Pt(2)

    def add_footer(slide, current_idx, total_slides=TOTAL_SLIDES):
        footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.12), Inches(11.733), Inches(0.32))
        tf = footer_box.text_frame
        tf.margin_left = Inches(0)
        tf.margin_top = Inches(0)
        p = tf.paragraphs[0]
        p.text = f"Đồ Án Chuyên Đề Back-End: Royal Bistro  |  Sinh viên: Nguyễn Hoàng Hưng – 501250384 (K25)                                   Trang {current_idx}/{total_slides}"
        p.font.name = "Segoe UI"
        p.font.size = Pt(10.5)
        p.font.color.rgb = RGBColor(148, 163, 184)

    # ============================================================
    # SLIDE 1: TRANG BÌA (HOÀN HẢO 3 THẺ CÂN ĐỐI, CHỮ TO RÕ)
    # ============================================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide1, LIGHT_BG)

    top_bar = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = VIBRANT_BLUE
    top_bar.line.fill.background()

    sch_box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(0.44), Inches(11.533), Inches(0.56))
    sch_box.adjustments[0] = 0.2
    sch_box.fill.solid()
    sch_box.fill.fore_color.rgb = WHITE
    sch_box.line.color.rgb = CARD_BORDER
    sch_box.line.width = Pt(1)
    sch_tf = sch_box.text_frame
    sch_tf.margin_top = Inches(0.1)
    sch_p = sch_tf.paragraphs[0]
    sch_p.text = "🏛️ TRƯỜNG CAO ĐẲNG CÔNG NGHỆ THÔNG TIN TP.HCM • KHOA CÔNG NGHỆ THÔNG TIN – ĐIỆN TỬ"
    sch_p.alignment = PP_ALIGN.CENTER
    sch_p.font.name = "Segoe UI"
    sch_p.font.size = Pt(14)
    sch_p.font.bold = True
    sch_p.font.color.rgb = PRIMARY

    main_card = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(1.12), Inches(11.533), Inches(2.72))
    main_card.adjustments[0] = 0.05
    main_card.fill.solid()
    main_card.fill.fore_color.rgb = WHITE
    main_card.line.color.rgb = RGBColor(199, 210, 254)
    main_card.line.width = Pt(1.5)
    
    m_tf = main_card.text_frame
    m_tf.word_wrap = True
    m_tf.margin_top = Inches(0.18)
    m_tf.margin_left = Inches(0.4)
    m_tf.margin_right = Inches(0.4)

    mp0 = m_tf.paragraphs[0]
    mp0.text = "BÁO CÁO ĐỒ ÁN MÔN HỌC: CHUYÊN ĐỀ BACK-END"
    mp0.alignment = PP_ALIGN.CENTER
    mp0.font.name = "Segoe UI"
    mp0.font.size = Pt(15.5)
    mp0.font.bold = True
    mp0.font.color.rgb = VIBRANT_BLUE

    mp1 = m_tf.add_paragraph()
    mp1.text = "XÂY DỰNG HỆ THỐNG BACK-END CHO PHẦN MỀM\nQUẢN LÝ NHÀ HÀNG & GỌI MÓN (ROYAL BISTRO)"
    mp1.alignment = PP_ALIGN.CENTER
    mp1.font.name = "Segoe UI"
    mp1.font.size = Pt(28.5)
    mp1.font.bold = True
    mp1.font.color.rgb = PRIMARY
    mp1.space_before = Pt(6)

    mp2 = m_tf.add_paragraph()
    mp2.text = "Kiến trúc RESTful API & WebSocket Realtime • Điều phối Bếp KDS • Tự động trừ kho định mức BOM • Báo cáo doanh thu"
    mp2.alignment = PP_ALIGN.CENTER
    mp2.font.name = "Segoe UI"
    mp2.font.size = Pt(13.5)
    mp2.font.color.rgb = TEXT_MUTED
    mp2.space_before = Pt(6)

    card_data = [
        ("👤 SINH VIÊN THỰC HIỆN", [
            ("Họ và tên:", "Nguyễn Hoàng Hưng"),
            ("MSSV:", "501250384"),
            ("Khóa đào tạo:", "Khóa K25"),
            ("Chuyên ngành:", "Công nghệ Phần mềm")
        ], VIBRANT_BLUE, RGBColor(238, 242, 255)),

        ("👨‍🏫 GIẢNG VIÊN HƯỚNG DẪN", [
            ("Giảng viên:", "ThS. NGUYỄN MINH HẢI"),
            ("Đơn vị:", "Khoa CNTT – ĐT (ITC)"),
            ("Môn học:", "Chuyên Đề Back-End"),
            ("Thời gian:", "Học kỳ I / 2026")
        ], EMERALD, RGBColor(236, 253, 245)),

        ("⚡ CÔNG NGHỆ CHÍNH", [
            ("Máy chủ:", "Node.js v20 & Express.js"),
            ("Cơ sở dữ liệu:", "MongoDB v7 & Mongoose"),
            ("Realtime:", "Socket.io WebSockets"),
            ("Bảo mật:", "JWT Stateless & Bcrypt")
        ], GOLD, RGBColor(254, 243, 199))
    ]

    c_left = 0.9
    for c_title, rows, c_color, c_bg in card_data:
        box = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left), Inches(3.98), Inches(3.68), Inches(2.82))
        box.adjustments[0] = 0.06
        box.fill.solid()
        box.fill.fore_color.rgb = WHITE
        box.line.color.rgb = c_color
        box.line.width = Pt(1.5)

        pill = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(c_left + 0.15), Inches(4.14), Inches(3.38), Inches(0.44))
        pill.adjustments[0] = 0.3
        pill.fill.solid()
        pill.fill.fore_color.rgb = c_bg
        pill.line.fill.background()
        ptf = pill.text_frame
        ptf.margin_top = Inches(0.08)
        pp = ptf.paragraphs[0]
        pp.text = c_title
        pp.alignment = PP_ALIGN.CENTER
        pp.font.name = "Segoe UI"
        pp.font.size = Pt(13.5)
        pp.font.bold = True
        pp.font.color.rgb = c_color

        tx = slide1.shapes.add_textbox(Inches(c_left + 0.16), Inches(4.68), Inches(3.36), Inches(1.98))
        itf = tx.text_frame
        itf.margin_top = Inches(0)
        itf.margin_left = Inches(0)
        for idx, (lbl, val) in enumerate(rows):
            p = itf.paragraphs[0] if idx == 0 else itf.add_paragraph()
            p.text = f"{lbl} "
            p.font.name = "Segoe UI"
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_MUTED
            if idx > 0:
                p.space_before = Pt(6)
            run = p.add_run()
            run.text = val
            run.font.bold = True
            run.font.color.rgb = TEXT_DARK

        c_left += 3.93

    add_footer(slide1, 1)

    # ============================================================
    # SLIDE 2: MỤC ĐÍCH DỰ ÁN & GIÁ TRỊ THỰC TIỄN (CHỮ TO RÕ, ĐỊNH TÍNH)
    # ============================================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide2, LIGHT_BG)
    add_header(slide2, "Mục Tiêu & Tác Động Thực Tế", "Mục Đích Dự Án & Giá Trị Thực Tiễn Cho Ngành F&B", "Khắc phục 3 bài toán nan giải của nhà hàng truyền thống và số hóa toàn diện quy trình vận hành", tag_width=4.4)
    add_footer(slide2, 2)

    # Khung trái: Thực trạng nan giải
    b_left = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(4.15))
    b_left.adjustments[0] = 0.05
    b_left.fill.solid()
    b_left.fill.fore_color.rgb = WHITE
    b_left.line.color.rgb = RGBColor(254, 202, 202)
    b_left.line.width = Pt(1.5)
    tf_l = b_left.text_frame
    tf_l.word_wrap = True
    tf_l.margin_top = Inches(0.18)
    tf_l.margin_left = Inches(0.25)
    tf_l.margin_right = Inches(0.25)

    p_lh = tf_l.paragraphs[0]
    p_lh.text = "⚠️ 3 BÀI TOÁN THỰC TẾ NAN GIẢI"
    p_lh.font.name = "Segoe UI"
    p_lh.font.size = Pt(15.5)
    p_lh.font.bold = True
    p_lh.font.color.rgb = RED

    pains = [
        ("Tắc nghẽn giờ cao điểm", "Khách chờ gọi món và thanh toán lâu trong giờ đông khách; nhân viên chạy bàn quá tải, dễ sót order."),
        ("Thất thoát kho nguyên liệu", "Đầu bếp nấu ước lượng tay; không có định mức (BOM) liên kết trực tiếp với kho gây hao hụt chi phí."),
        ("Sai sót & chậm trễ order", "Ghi chép giấy dễ nhầm món, sót ghi chú của khách; truyền đạt từ bàn vào bếp bị gián đoạn.")
    ]
    for title, desc in pains:
        p_item = tf_l.add_paragraph()
        p_item.text = f"• {title}: "
        p_item.font.name = "Segoe UI"
        p_item.font.size = Pt(13.5)
        p_item.font.bold = True
        p_item.font.color.rgb = TEXT_DARK
        p_item.space_before = Pt(10)
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Khung phải: Giá trị thực tiễn
    b_right = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.35), Inches(5.7), Inches(4.15))
    b_right.adjustments[0] = 0.05
    b_right.fill.solid()
    b_right.fill.fore_color.rgb = WHITE
    b_right.line.color.rgb = RGBColor(167, 243, 208)
    b_right.line.width = Pt(1.5)
    tf_r = b_right.text_frame
    tf_r.word_wrap = True
    tf_r.margin_top = Inches(0.18)
    tf_r.margin_left = Inches(0.25)
    tf_r.margin_right = Inches(0.25)

    p_rh = tf_r.paragraphs[0]
    p_rh.text = "💎 3 GIÁ TRỊ THỰC TIỄN VƯỢT TRỘI"
    p_rh.font.name = "Segoe UI"
    p_rh.font.size = Pt(15.5)
    p_rh.font.bold = True
    p_rh.font.color.rgb = EMERALD

    goals = [
        ("Tối ưu nhân sự phục vụ", "Khách tự quét mã QR duyệt thực đơn và gọi món tại bàn trên điện thoại; tinh gọn số lượng nhân viên chạy bàn."),
        ("Kiểm soát kho chính xác", "Thuật toán tự động trừ kho theo định mức BOM khi bếp bấm nấu; cảnh báo kịp thời khi nguyên liệu sắp cạn."),
        ("Truyền đơn tức thời vào Bếp", "Đơn gọi món gửi trực tiếp vào màn hình Bếp KDS qua WebSocket; chuông reo nhận đơn để chế biến ngay lập tức.")
    ]
    for title, desc in goals:
        p_item = tf_r.add_paragraph()
        p_item.text = f"• {title}: "
        p_item.font.name = "Segoe UI"
        p_item.font.size = Pt(13.5)
        p_item.font.bold = True
        p_item.font.color.rgb = TEXT_DARK
        p_item.space_before = Pt(10)
        run = p_item.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # 3 Hộp giá trị thực tiễn nổi bật
    metrics = [
        ("⚡ Real-Time", "Truyền đơn vào Bếp tức thời", VIBRANT_BLUE),
        ("👥 Quét Mã QR", "Khách chủ động gọi món tại bàn", SKY_BLUE),
        ("📦 Định Mức BOM", "Tự động trừ kho theo từng món", EMERALD)
    ]
    m_left = 0.8
    for m_val, m_desc, m_col in metrics:
        m_box = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(m_left), Inches(5.65), Inches(3.75), Inches(1.28))
        m_box.adjustments[0] = 0.08
        m_box.fill.solid()
        m_box.fill.fore_color.rgb = WHITE
        m_box.line.color.rgb = CARD_BORDER
        m_box.line.width = Pt(1.5)
        m_tf = m_box.text_frame
        m_tf.word_wrap = True
        m_tf.margin_top = Inches(0.18)
        m_tf.margin_left = Inches(0.15)
        m_tf.margin_right = Inches(0.15)
        
        mp1 = m_tf.paragraphs[0]
        mp1.text = m_val
        mp1.alignment = PP_ALIGN.CENTER
        mp1.font.name = "Segoe UI"
        mp1.font.size = Pt(24)
        mp1.font.bold = True
        mp1.font.color.rgb = m_col

        mp2 = m_tf.add_paragraph()
        mp2.text = m_desc
        mp2.alignment = PP_ALIGN.CENTER
        mp2.font.name = "Segoe UI"
        mp2.font.size = Pt(13)
        mp2.font.bold = True
        mp2.font.color.rgb = TEXT_MUTED
        mp2.space_before = Pt(3)
        m_left += 3.99

    # ============================================================
    # SLIDE 3: BỘ CÔNG CỤ SỬ DỤNG ĐỂ BUILD (CHỮ TO RÕ)
    # ============================================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide3, LIGHT_BG)
    add_header(slide3, "Công Cụ & Hệ Sinh Thái Build", "Bộ Công Cụ & Công Nghệ Sử Dụng Để Xây Dựng Hệ Thống", "Bộ công nghệ tuyển chọn hiện đại: Non-blocking I/O, thời gian thực và tự động hóa kiểm thử", tag_width=4.5)
    add_footer(slide3, 3)

    tool_blocks = [
        ("🖥️ NỀN TẢNG MÁY CHỦ (SERVER RUNTIME)", [
            ("Node.js v20 LTS:", "Môi trường thực thi JavaScript Non-blocking I/O, xử lý đa luồng bất đồng bộ tải cao."),
            ("Express.js v4.19:", "Framework Back-End phân tầng MVC, quản lý pipeline middleware và định tuyến RESTful API.")
        ], VIBRANT_BLUE, RGBColor(238, 242, 255)),

        ("🍃 CƠ SỞ DỮ LIỆU & ORM (DATABASE)", [
            ("MongoDB v7.0:", "Hệ quản trị CSDL NoSQL hướng tài liệu (JSON-like), mở rộng linh hoạt, lưu trữ 12 Collections."),
            ("Mongoose v8.3 ODM:", "Mô hình hóa Schema chặt chẽ, ràng buộc dữ liệu toàn vẹn và thực thi Aggregation Pipeline.")
        ], EMERALD, RGBColor(236, 253, 245)),

        ("⚡ KÊNH REAL-TIME & BẢO MẬT", [
            ("Socket.io v4.7:", "Kênh WebSocket 2 chiều phát sóng sự kiện thời gian thực giữa Bàn khách, POS và Bếp KDS."),
            ("JWT & Bcrypt:", "Xác thực Token không lưu trạng thái (Stateless 24h) và băm mật khẩu bảo mật một chiều 10 salt rounds.")
        ], GOLD, RGBColor(254, 243, 199)),

        ("🛠️ CÔNG CỤ PHÁT TRIỂN & BUILD (DEV TOOLS)", [
            ("VS Code & Postman:", "Môi trường lập trình chính và công cụ gửi request mô phỏng kiểm thử tự động toàn diện các API endpoints."),
            ("Git / GitHub & NPM:", "Quản lý phiên bản mã nguồn phân tán Monorepo, quản lý gói phụ thuộc và tự động hóa build.")
        ], SKY_BLUE, RGBColor(240, 249, 255))
    ]

    coords = [
        (0.8, 1.4),   # Top-left
        (6.8, 1.4),   # Top-right
        (0.8, 4.25),  # Bottom-left
        (6.8, 4.25)   # Bottom-right
    ]

    for (bx, by), (b_title, b_rows, b_col, b_bg) in zip(coords, tool_blocks):
        tb_box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(bx), Inches(by), Inches(5.7), Inches(2.65))
        tb_box.adjustments[0] = 0.05
        tb_box.fill.solid()
        tb_box.fill.fore_color.rgb = WHITE
        tb_box.line.color.rgb = b_col
        tb_box.line.width = Pt(1.5)

        tb_tf = tb_box.text_frame
        tb_tf.word_wrap = True
        tb_tf.margin_top = Inches(0.18)
        tb_tf.margin_left = Inches(0.25)
        tb_tf.margin_right = Inches(0.25)

        p0 = tb_tf.paragraphs[0]
        p0.text = b_title
        p0.font.name = "Segoe UI"
        p0.font.size = Pt(14.5)
        p0.font.bold = True
        p0.font.color.rgb = b_col

        for row_lbl, row_desc in b_rows:
            p = tb_tf.add_paragraph()
            p.text = f"• {row_lbl} "
            p.font.name = "Segoe UI"
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = TEXT_DARK
            p.space_before = Pt(8)
            run = p.add_run()
            run.text = row_desc
            run.font.bold = False
            run.font.color.rgb = TEXT_MUTED

    # ============================================================
    # SLIDE 4: KIẾN TRÚC MÁY CHỦ & LUỒNG XỬ LÝ DỮ LIỆU
    # ============================================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide4, LIGHT_BG)
    add_header(slide4, "Kiến Trúc & Tổ Chức Hệ Thống", "Kiến Trúc Máy Chủ Client-Server & Phân Tầng Mã Nguồn Back-End", "Mô hình phân tầng rõ ràng, kết hợp RESTful API với Kênh sự kiện thời gian thực WebSocket", tag_width=4.5)
    add_footer(slide4, 4)

    # Cột Trái: Cấu trúc thư mục Monorepo
    left_m = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.7), Inches(5.55))
    left_m.adjustments[0] = 0.05
    left_m.fill.solid()
    left_m.fill.fore_color.rgb = WHITE
    left_m.line.color.rgb = CARD_BORDER
    left_m.line.width = Pt(1.5)
    tf_m = left_m.text_frame
    tf_m.word_wrap = True
    tf_m.margin_top = Inches(0.2)
    tf_m.margin_left = Inches(0.25)
    tf_m.margin_right = Inches(0.25)

    pm_h = tf_m.paragraphs[0]
    pm_h.text = "📁 PHÂN TẦNG MÃ NGUỒN BACK-END (BE/src)"
    pm_h.font.name = "Segoe UI"
    pm_h.font.size = Pt(15.5)
    pm_h.font.bold = True
    pm_h.font.color.rgb = PRIMARY

    dirs = [
        ("controllers/", "Xử lý nghiệp vụ đặt bàn, gọi món, bếp KDS, kho BOM và tính doanh thu."),
        ("middlewares/", "Xác thực JWT token, kiểm tra phân quyền RBAC 4 vai trò, bắt lỗi tập trung."),
        ("models/", "12 Lược đồ Mongoose Schema & ràng buộc quan hệ n-n (BOM)."),
        ("routes/", "Định tuyến hơn 30 RESTful API endpoints theo chuẩn RESTful conventions."),
        ("utils/ & config/", "Kết nối CSDL MongoDB, tiện ích phát sóng Socket.io và xử lý asyncHandler."),
        ("server.js", "Entry point khởi động Express HTTP Server & Socket.io WebSocket.")
    ]
    for d_name, d_desc in dirs:
        pd = tf_m.add_paragraph()
        pd.text = f"• {d_name} "
        pd.font.name = "Segoe UI"
        pd.font.size = Pt(13)
        pd.font.bold = True
        pd.font.color.rgb = VIBRANT_BLUE
        pd.space_before = Pt(7)
        run = pd.add_run()
        run.text = d_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Cột Phải: Luồng xử lý request 5 bước trực quan
    right_m = slide4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(1.35), Inches(5.7), Inches(5.55))
    right_m.adjustments[0] = 0.05
    right_m.fill.solid()
    right_m.fill.fore_color.rgb = WHITE
    right_m.line.color.rgb = CARD_BORDER
    right_m.line.width = Pt(1.5)
    tf_r4 = right_m.text_frame
    tf_r4.word_wrap = True
    tf_r4.margin_top = Inches(0.2)
    tf_r4.margin_left = Inches(0.25)
    tf_r4.margin_right = Inches(0.25)

    pr_h = tf_r4.paragraphs[0]
    pr_h.text = "🔄 LUỒNG XỬ LÝ REQUEST & REALTIME EVENT"
    pr_h.font.name = "Segoe UI"
    pr_h.font.size = Pt(15.5)
    pr_h.font.bold = True
    pr_h.font.color.rgb = EMERALD

    steps = [
        ("1. Khởi tạo yêu cầu:", "Khách quét QR hoặc Thu ngân POS gửi HTTP Request / Socket Event lên máy chủ."),
        ("2. Gateway bảo mật:", "Middleware (Helmet, CORS, Rate Limit) chặn tấn công; xác thực JWT Token và quyền vai trò."),
        ("3. Xử lý Controller:", "Controller kiểm tra tính hợp lệ dữ liệu, điều phối logic nghiệp vụ phù hợp."),
        ("4. Tương tác CSDL:", "Mongoose ODM thực thi truy vấn MongoDB (CRUD, Transaction, Aggregation Pipeline)."),
        ("5. Đồng bộ Realtime:", "Kích hoạt Socket.io phát sóng tức thời xuống màn hình Bếp KDS và Bàn ăn.")
    ]
    for s_title, s_desc in steps:
        ps = tf_r4.add_paragraph()
        ps.text = f"• {s_title} "
        ps.font.name = "Segoe UI"
        ps.font.size = Pt(13)
        ps.font.bold = True
        ps.font.color.rgb = TEXT_DARK
        ps.space_before = Pt(7)
        run = ps.add_run()
        run.text = s_desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # ============================================================
    # SLIDE 5: BẢN ĐỒ TỔNG QUAN 9 TÍNH NĂNG NGHIỆP VỤ (BẢNG RÕ RÀNG)
    # ============================================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide5, LIGHT_BG)
    add_header(slide5, "Bản Đồ Tính Năng Hệ Thống", "Tổng Quan 9 Phân Hệ Tính Năng Nghiệp Vụ Cốt Lõi Của Back-End", "Bao phủ toàn diện vòng đời vận hành nhà hàng: Xác thực, Bàn ăn, Gọi món, Bếp KDS, Kho BOM, CRM và Báo cáo", tag_width=4.2)
    add_footer(slide5, 5)

    rows = 10
    cols = 3
    table_shape = slide5.shapes.add_table(rows, cols, Inches(0.8), Inches(1.35), Inches(11.733), Inches(5.55))
    tbl = table_shape.table
    tbl.columns[0].width = Inches(3.2)
    tbl.columns[1].width = Inches(2.3)
    tbl.columns[2].width = Inches(6.233)

    headers = ["Phân Hệ Nghiệp Vụ", "Tuyến Đường API (Route)", "Mô Tả Chức Năng Trọng Tâm Của Back-End"]
    for c_idx, h_text in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PRIMARY
        cp = cell.text_frame.paragraphs[0]
        cp.text = h_text
        cp.font.name = "Segoe UI"
        cp.font.size = Pt(12.5)
        cp.font.bold = True
        cp.font.color.rgb = WHITE
        cell.text_frame.margin_top = Inches(0.06)
        cell.text_frame.margin_bottom = Inches(0.06)

    mod_data = [
        ("1. Xác Thực & Phân Quyền RBAC", "/api/v1/auth", "Đăng nhập, cấp JWT 24h, băm Bcrypt 10 rounds, phân quyền 4 vai trò (Admin, Thu ngân, Bếp, Phục vụ)"),
        ("2. Quản Lý Sơ Đồ Bàn Ăn", "/api/v1/tables", "Giám sát trạng thái bàn động (Trống, Có khách, Đã đặt), chuyển bàn, gộp bàn và tự động giải phóng bàn"),
        ("3. Thực Đơn & Danh Mục Món", "/api/v1/dishes", "Phân loại món ăn (Khai vị, Món chính, Đồ uống...), giá niêm yết, bật/tắt phục vụ và upload ảnh đám mây"),
        ("4. Điểm Bán POS & Khách Quét QR", "/api/v1/orders", "Lập đơn gọi món tại bàn, ghi chú khẩu vị (ít cay, không hành...), phát sóng Socket.io sự kiện đơn mới tức thời"),
        ("5. Điều Phối Màn Hình Bếp KDS", "/api/v1/orders/:id/status", "Tiếp nhận đơn tức thời, chuông reo báo hiệu, cập nhật tiến độ nấu 1 chạm và kích hoạt khấu trừ kho BOM"),
        ("6. Kho Nguyên Liệu & Định Mức BOM", "/api/v1/inventory", "Thuật toán tự động trừ kho theo công thức BOM từng gram, cảnh báo cạn kiệt khi chạm ngưỡng tồn tối thiểu"),
        ("7. Đặt Bàn Tiệc Trước (Reservation)", "/api/v1/reservations", "Khách hẹn đặt bàn trước, kiểm tra bàn trống theo giờ, nhận tiền cọc giữ chỗ và tự động check-in khi khách đến"),
        ("8. Khách Hàng Thân Thiết & CRM", "/api/v1/customers", "Tra cứu SĐT khách hàng, tự động tích điểm thưởng (100k = 1 điểm), nâng hạng thẻ VIP và chiết khấu trực tiếp"),
        ("9. Báo Cáo Doanh Thu & Thống Kê", "/api/v1/reports", "MongoDB Aggregation tính doanh thu theo ngày/tuần/tháng, tỷ lệ dòng tiền VietQR/tiền mặt, Top 5 món bán chạy")
    ]
    for r_idx, (m_name, m_route, m_func) in enumerate(mod_data, start=1):
        bg_col = WHITE if r_idx % 2 != 0 else CARD_ALT
        for c_idx, text_val in enumerate([m_name, m_route, m_func]):
            cell = tbl.cell(r_idx, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = bg_col
            cp = cell.text_frame.paragraphs[0]
            cp.text = text_val
            cp.font.name = "Segoe UI"
            cp.font.size = Pt(11.5)
            if c_idx == 0:
                cp.font.bold = True
                cp.font.color.rgb = TEXT_DARK
            elif c_idx == 1:
                cp.font.color.rgb = VIBRANT_BLUE
                cp.font.name = "Consolas"
                cp.font.size = Pt(11)
            else:
                cp.font.color.rgb = TEXT_MUTED
            cell.text_frame.margin_top = Inches(0.04)
            cell.text_frame.margin_bottom = Inches(0.04)

    # ============================================================
    # SLIDE 6: TÍNH NĂNG CỐT LÕI (PHẦN 1: BÀN ĂN & GỌI MÓN)
    # ============================================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide6, LIGHT_BG)
    add_header(slide6, "Tính Năng Thực Nghiệm (1 & 2)", "Sơ Đồ Bàn Ăn Thời Gian Thực & Gọi Món Đa Kênh (POS / Khách QR)", "Số hóa toàn bộ khâu tiếp đón khách, mở bàn và truyền nhận đơn hàng siêu tốc", tag_width=4.5)
    add_footer(slide6, 6)

    left_f6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.5), Inches(5.55))
    left_f6.adjustments[0] = 0.05
    left_f6.fill.solid()
    left_f6.fill.fore_color.rgb = WHITE
    left_f6.line.color.rgb = CARD_BORDER
    left_f6.line.width = Pt(1.5)
    tf_f6 = left_f6.text_frame
    tf_f6.word_wrap = True
    tf_f6.margin_top = Inches(0.2)
    tf_f6.margin_left = Inches(0.25)
    tf_f6.margin_right = Inches(0.25)

    p6_h1 = tf_f6.paragraphs[0]
    p6_h1.text = "🪑 1. SƠ ĐỒ BÀN ĂN REALTIME"
    p6_h1.font.name = "Segoe UI"
    p6_h1.font.size = Pt(15.5)
    p6_h1.font.bold = True
    p6_h1.font.color.rgb = PRIMARY

    f6_items1 = [
        ("Trạng thái động 3 màu:", "Bàn trống (Xanh lá), Có khách (Đỏ), Khách đã đặt hẹn trước (Vàng cam)."),
        ("Tự động giải phóng bàn:", "Khi thu ngân xác nhận thanh toán ➔ Hệ thống tự động trả bàn về trạng thái 'trống' tức thì.")
    ]
    for lbl, desc in f6_items1:
        p = tf_f6.add_paragraph()
        p.text = f"• {lbl} "
        p.font.name = "Segoe UI"
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = EMERALD
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    p6_h2 = tf_f6.add_paragraph()
    p6_h2.text = "📱 2. GỌI MÓN ĐA KÊNH & BẮN ĐƠN BẾP"
    p6_h2.font.name = "Segoe UI"
    p6_h2.font.size = Pt(15.5)
    p6_h2.font.bold = True
    p6_h2.font.color.rgb = PRIMARY
    p6_h2.space_before = Pt(16)

    f6_items2 = [
        ("Khách tự quét mã QR:", "Khách tự gọi món trên điện thoại, tùy chỉnh khẩu vị (ít cay, không hành, ít đá...)."),
        ("POS Phục vụ bàn:", "Nhân viên mở bàn, thêm món nhanh chóng bằng màn hình cảm ứng POS."),
        ("Bắn đơn Socket.io tức thời:", "Xác nhận gọi món ➔ Sự kiện 'order:new' bắn thẳng xuống Bếp KDS chuông reo tức thì.")
    ]
    for lbl, desc in f6_items2:
        p = tf_f6.add_paragraph()
        p.text = f"• {lbl} "
        p.font.name = "Segoe UI"
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = VIBRANT_BLUE
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Cột phải: Ảnh chụp màn hình Sơ đồ bàn ăn sạch sẽ
    r_box6 = slide6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.35), Inches(5.933), Inches(5.55))
    r_box6.adjustments[0] = 0.05
    r_box6.fill.solid()
    r_box6.fill.fore_color.rgb = WHITE
    r_box6.line.color.rgb = CARD_BORDER
    r_box6.line.width = Pt(1.5)

    r_cap6 = slide6.shapes.add_textbox(Inches(6.8), Inches(1.48), Inches(5.5), Inches(0.38))
    rc_tf6 = r_cap6.text_frame
    rc_p6 = rc_tf6.paragraphs[0]
    rc_p6.text = "📸 Giao diện Sơ Đồ Bàn Ăn Thời Gian Thực (Table Map)"
    rc_p6.font.name = "Segoe UI"
    rc_p6.font.size = Pt(13.5)
    rc_p6.font.bold = True
    rc_p6.font.color.rgb = PRIMARY

    img_tables = os.path.join(clean_screenshots_dir, "04_tables.png")
    if os.path.exists(img_tables):
        slide6.shapes.add_picture(img_tables, Inches(6.75), Inches(1.95), width=Inches(5.633))

    # ============================================================
    # SLIDE 7: TÍNH NĂNG CỐT LÕI (PHẦN 2: BẾP KDS, KHO BOM & BÁO CÁO)
    # ============================================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide7, LIGHT_BG)
    add_header(slide7, "Tính Năng Thực Nghiệm (3 & 4)", "Điều Phối Bếp KDS, Trừ Kho BOM & Báo Cáo Doanh Thu", "Tự động hóa khấu trừ nguyên liệu từng gram và phân tích tài chính thông minh", tag_width=4.5)
    add_footer(slide7, 7)

    left_f7 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.35), Inches(5.5), Inches(5.55))
    left_f7.adjustments[0] = 0.05
    left_f7.fill.solid()
    left_f7.fill.fore_color.rgb = WHITE
    left_f7.line.color.rgb = CARD_BORDER
    left_f7.line.width = Pt(1.5)
    tf_f7 = left_f7.text_frame
    tf_f7.word_wrap = True
    tf_f7.margin_top = Inches(0.2)
    tf_f7.margin_left = Inches(0.25)
    tf_f7.margin_right = Inches(0.25)

    p7_h1 = tf_f7.paragraphs[0]
    p7_h1.text = "🍳 3. ĐIỀU PHỐI BẾP KDS & TRỪ KHO BOM"
    p7_h1.font.name = "Segoe UI"
    p7_h1.font.size = Pt(15.5)
    p7_h1.font.bold = True
    p7_h1.font.color.rgb = PRIMARY

    f7_items1 = [
        ("Màn hình Bếp KDS 1 chạm:", "Chuông reo báo đơn; đầu bếp bấm 'Bắt Đầu Nấu' ➔ 'Hoàn Tất' để đồng bộ tiến độ món."),
        ("Thuật toán tự động trừ kho:", "Bếp bấm nấu ➔ Back-End tự tính toán theo bảng định mức BOM và trừ thẳng kho nguyên liệu từng gram.")
    ]
    for lbl, desc in f7_items1:
        p = tf_f7.add_paragraph()
        p.text = f"• {lbl} "
        p.font.name = "Segoe UI"
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = GOLD
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    p7_h2 = tf_f7.add_paragraph()
    p7_h2.text = "📊 4. BÁO CÁO DOANH THU & DÒNG TIỀN"
    p7_h2.font.name = "Segoe UI"
    p7_h2.font.size = Pt(15.5)
    p7_h2.font.bold = True
    p7_h2.font.color.rgb = PRIMARY
    p7_h2.space_before = Pt(16)

    f7_items2 = [
        ("MongoDB Aggregation Pipeline:", "Truy vấn đường ống tính toán doanh số tự động theo thời gian thực."),
        ("Phân tích cơ cấu dòng tiền:", "Thống kê tỷ lệ thanh toán chuyển khoản VietQR vs. Tiền mặt; xếp hạng Top món bán chạy nhất.")
    ]
    for lbl, desc in f7_items2:
        p = tf_f7.add_paragraph()
        p.text = f"• {lbl} "
        p.font.name = "Segoe UI"
        p.font.size = Pt(13.5)
        p.font.bold = True
        p.font.color.rgb = EMERALD
        p.space_before = Pt(8)
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED

    # Cột phải: Ảnh chụp màn hình Bếp KDS sạch sẽ
    r_box7 = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.6), Inches(1.35), Inches(5.933), Inches(5.55))
    r_box7.adjustments[0] = 0.05
    r_box7.fill.solid()
    r_box7.fill.fore_color.rgb = WHITE
    r_box7.line.color.rgb = CARD_BORDER
    r_box7.line.width = Pt(1.5)

    r_cap7 = slide7.shapes.add_textbox(Inches(6.8), Inches(1.48), Inches(5.5), Inches(0.38))
    rc_tf7 = r_cap7.text_frame
    rc_p7 = rc_tf7.paragraphs[0]
    rc_p7.text = "📸 Giao diện Điều Phối Màn Hình Bếp KDS Thông Minh"
    rc_p7.font.name = "Segoe UI"
    rc_p7.font.size = Pt(13.5)
    rc_p7.font.bold = True
    rc_p7.font.color.rgb = PRIMARY

    img_kds = os.path.join(clean_screenshots_dir, "05_kitchen.png")
    if os.path.exists(img_kds):
        slide7.shapes.add_picture(img_kds, Inches(6.75), Inches(1.95), width=Inches(5.633))

    # ============================================================
    # SLIDE 8: CẢM ƠN & PHẦN HỎI ĐÁP (CHUẨN HOÀN TOÀN, CHỮ TO RÕ)
    # ============================================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_bg(slide8, LIGHT_BG)

    top_bar8 = slide8.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(0.12))
    top_bar8.fill.solid()
    top_bar8.fill.fore_color.rgb = VIBRANT_BLUE
    top_bar8.line.fill.background()

    # Khung chính giữa to, sang trọng
    thank_card = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(0.85), Inches(10.933), Inches(4.3))
    thank_card.adjustments[0] = 0.05
    thank_card.fill.solid()
    thank_card.fill.fore_color.rgb = WHITE
    thank_card.line.color.rgb = RGBColor(199, 210, 254)
    thank_card.line.width = Pt(1.5)

    t_tf = thank_card.text_frame
    t_tf.word_wrap = True
    t_tf.margin_top = Inches(0.32)
    t_tf.margin_left = Inches(0.5)
    t_tf.margin_right = Inches(0.5)

    tp0 = t_tf.paragraphs[0]
    tp0.text = "🏛️ KHOA CÔNG NGHỆ THÔNG TIN – ĐIỆN TỬ • ITC"
    tp0.alignment = PP_ALIGN.CENTER
    tp0.font.name = "Segoe UI"
    tp0.font.size = Pt(15)
    tp0.font.bold = True
    tp0.font.color.rgb = VIBRANT_BLUE

    tp1 = t_tf.add_paragraph()
    tp1.text = "XIN CHÂN THÀNH CẢM ƠN!"
    tp1.alignment = PP_ALIGN.CENTER
    tp1.font.name = "Segoe UI"
    tp1.font.size = Pt(34)
    tp1.font.bold = True
    tp1.font.color.rgb = PRIMARY
    tp1.space_before = Pt(12)

    tp2 = t_tf.add_paragraph()
    tp2.text = "Kính chúc Quý Thầy Cô trong Hội đồng thật nhiều sức khỏe, hạnh phúc và thành công!"
    tp2.alignment = PP_ALIGN.CENTER
    tp2.font.name = "Segoe UI"
    tp2.font.size = Pt(16.5)
    tp2.font.color.rgb = TEXT_MUTED
    tp2.space_before = Pt(8)

    tp3 = t_tf.add_paragraph()
    tp3.text = "❓ PHẦN HỎI & ĐÁP (Q & A)"
    tp3.alignment = PP_ALIGN.CENTER
    tp3.font.name = "Segoe UI"
    tp3.font.size = Pt(23)
    tp3.font.bold = True
    tp3.font.color.rgb = GOLD
    tp3.space_before = Pt(18)

    tp4 = t_tf.add_paragraph()
    tp4.text = "Em rất mong nhận được những câu hỏi phản biện, lời nhận xét và góp ý quý báu của Quý Thầy Cô để hoàn thiện đề tài."
    tp4.alignment = PP_ALIGN.CENTER
    tp4.font.name = "Segoe UI"
    tp4.font.size = Pt(15.5)
    tp4.font.color.rgb = TEXT_MUTED
    tp4.space_before = Pt(6)

    # Thẻ thông tin sinh viên ở đáy
    info_card8 = slide8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(5.35), Inches(10.933), Inches(1.4))
    info_card8.adjustments[0] = 0.1
    info_card8.fill.solid()
    info_card8.fill.fore_color.rgb = RGBColor(238, 242, 255)
    info_card8.line.color.rgb = RGBColor(199, 210, 254)
    info_card8.line.width = Pt(1)

    i_tf8 = info_card8.text_frame
    i_tf8.word_wrap = True
    i_tf8.margin_top = Inches(0.2)
    i_tf8.margin_left = Inches(0.4)
    i_tf8.margin_right = Inches(0.4)

    ip1 = i_tf8.paragraphs[0]
    ip1.text = "👤 Sinh viên: NGUYỄN HOÀNG HƯNG   |   MSSV: 501250384   |   Lớp: Khóa K25 – Công nghệ Phần mềm"
    ip1.alignment = PP_ALIGN.CENTER
    ip1.font.name = "Segoe UI"
    ip1.font.size = Pt(14)
    ip1.font.bold = True
    ip1.font.color.rgb = PRIMARY

    ip2 = i_tf8.add_paragraph()
    ip2.text = "👨‍🏫 Giảng viên hướng dẫn: ThS. NGUYỄN MINH HẢI   |   Đề tài: Hệ Thống Back-End Quản Lý Nhà Hàng & Gọi Món (Royal Bistro)"
    ip2.alignment = PP_ALIGN.CENTER
    ip2.font.name = "Segoe UI"
    ip2.font.size = Pt(13)
    ip2.font.color.rgb = TEXT_MUTED
    ip2.space_before = Pt(6)

    add_footer(slide8, 8, total_slides=8)

    # ============================================================
    # LƯU FILE POWERPOINT & ĐỒNG BỘ
    # ============================================================
    output_path = os.path.join(base_dir, "BaoCao_TongQuan_DoAn_BackEnd_RoyalBistro.pptx")
    prs.save(output_path)
    print(f"✅ Đã tạo thành công PowerPoint tại: {output_path}")

    sync_targets = [
        "05_NguyenHoangHung.pptx",
        "05_NguyenNgocHaThao.pptx",
        "5_NguyenNgocHaThao.pptx"
    ]
    for target in sync_targets:
        t_path = os.path.join(base_dir, target)
        prs.save(t_path)
        print(f"✅ Đã đồng bộ sang: {t_path}")

if __name__ == "__main__":
    create_presentation()
