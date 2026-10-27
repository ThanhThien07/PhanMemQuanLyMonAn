import os
import docx
import xml.etree.ElementTree as ET

DOC_PATH = 'docs/bao_cao/501250384_NguyenHoangHung_CDFrontEnd.docx'

print(f"File size: {os.path.getsize(DOC_PATH):,} bytes")
doc = docx.Document(DOC_PATH)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total sections: {len(doc.sections)}")

for i, sec in enumerate(doc.sections):
    print(f"\n--- Section {i} ---")
    hdr_txt = [p.text for p in sec.header.paragraphs if p.text]
    ftr_txt = [p.text for p in sec.footer.paragraphs if p.text]
    print(f"  Header: {hdr_txt}")
    print(f"  Footer text: {ftr_txt}")
    xml_str = ET.tostring(sec._sectPr).decode('utf-8')
    has_pg_start = 'w:start="1"' in xml_str
    print(f"  Has w:start=\"1\": {has_pg_start}")

# Check images embedded
doc_zip_media = []
import zipfile
with zipfile.ZipFile(DOC_PATH) as z:
    for name in z.namelist():
        if name.startswith('word/media/'):
            doc_zip_media.append(name)
print(f"\nEmbedded media images: {len(doc_zip_media)} files")
for m in doc_zip_media:
    print(f"  {m}")

# Check key titles in paragraphs
print("\n--- Key Headings Found ---")
for p in doc.paragraphs:
    if p.text.startswith("CHƯƠNG") or p.text in ["LỜI MỞ ĐẦU", "LỜI CẢM ƠN", "MỤC LỤC", "DANH MỤC HÌNH ẢNH & BẢNG BIỂU", "TÀI LIỆU THAM KHẢO"]:
        print(f"  • {p.text}")
