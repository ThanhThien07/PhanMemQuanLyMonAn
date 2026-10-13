import os, sys, docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

sys.stdout.reconfigure(encoding='utf-8')

def build_report():
    print("Loading base template 501250121_NguyenQuangSang_CDBackEnd.docx...")
    doc = docx.Document('501250121_NguyenQuangSang_CDBackEnd.docx')
    body = doc._body._element

    # Update Cover Page (Section 0)
    # Paragraph 11 is title
    p11 = doc.paragraphs[11]
    p11.text = ""
    r11 = p11.add_run("ĐỀ TÀI: XÂY DỰNG HỆ THỐNG BACK-END CHO PHẦN MỀM QUẢN LÝ NHÀ HÀNG & GỌI MÓN (ROYAL BISTRO)")
    r11.font.name = "Times New Roman"
    r11.font.size = Pt(16)
    r11.font.bold = True
    p11.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Paragraph 29 is Giang vien huong dan
    p29 = doc.paragraphs[29]
    p29.text = ""
    r29_1 = p29.add_run("Giáo viên hướng dẫn :")
    r29_1.font.name = "Times New Roman"
    r29_1.font.size = Pt(13)
    r29_2 = p29.add_run(" ThS. Nguyễn Minh Hải")
    r29_2.font.name = "Times New Roman"
    r29_2.font.size = Pt(13)
    r29_2.font.bold = True

    # Paragraph 30 is Sinh vien thuc hien
    p30 = doc.paragraphs[30]
    p30.text = ""
    r30_1 = p30.add_run("Sinh viên thực hiện: ")
    r30_1.font.name = "Times New Roman"
    r30_1.font.size = Pt(13)
    r30_2 = p30.add_run("Nguyễn Ngọc Hưng – 501250384")
    r30_2.font.name = "Times New Roman"
    r30_2.font.size = Pt(13)
    r30_2.font.bold = True

    # Paragraph 31 is Khoa
    p31 = doc.paragraphs[31]
    p31.text = ""
    r31_1 = p31.add_run("Khoá\t")
    r31_1.font.name = "Times New Roman"
    r31_1.font.size = Pt(13)
    r31_2 = p31.add_run(": K25")
    r31_2.font.name = "Times New Roman"
    r31_2.font.size = Pt(13)
    r31_2.font.bold = True

    # Paragraph 34 is Location & Date
    p34 = doc.paragraphs[34]
    p34.text = ""
    r34 = p34.add_run("TP. Hồ Chí Minh, ngày 16 tháng 10 năm 2025")
    r34.font.name = "Times New Roman"
    r34.font.size = Pt(13)
    r34.font.bold = True

    # Truncate all elements after paragraph 34
    p34_idx = body.index(p34._p)
    for child in list(body)[p34_idx+1:]:
        body.remove(child)

    print("Cover page updated, body trimmed.")

    # Helper styling functions
    def format_run(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color=None):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = color

    def add_title(text):
        p = doc.add_paragraph()
        p.style = 'normal'
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(14)
        r = p.add_run(text)
        format_run(r, size_pt=14, bold=True)
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.style = 'Heading 1'
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        format_run(r, size_pt=16, bold=True)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.style = 'Heading 2'
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(11)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        format_run(r, size_pt=14, bold=True)
        return p

    def add_h3(text):
        p = doc.add_paragraph()
        p.style = 'Heading 3'
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(7)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        format_run(r, size_pt=12, bold=True)
        return p

    def add_p(text, bold_prefix=None, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = doc.add_paragraph()
        p.style = 'normal'
        p.alignment = align
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(0)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            format_run(r_pre, size_pt=12, bold=True)
        r = p.add_run(text)
        format_run(r, size_pt=12, italic=italic)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph()
        p.style = 'normal'
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(3)
        r_bullet = p.add_run("• ")
        format_run(r_bullet, size_pt=12, bold=True)
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            format_run(r_pre, size_pt=12, bold=True)
        r = p.add_run(text)
        format_run(r, size_pt=12)
        return p

    def add_fig_caption(text):
        p = doc.add_paragraph()
        p.style = 'Heading 4'
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = False
        r = p.add_run(text)
        format_run(r, size_pt=12, bold=True)
        return p

    def add_tbl_caption(text):
        p = doc.add_paragraph()
        p.style = 'Heading 5'
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        format_run(r, size_pt=12, bold=True)
        return p

    def add_fig(img_path, caption_text, width=Inches(5.5)):
        if os.path.exists(img_path):
            p = doc.add_paragraph()
            p.style = 'normal'
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            r = p.add_run()
            r.add_picture(img_path, width=width)
            add_fig_caption(caption_text)
        else:
            print(f"Warning: Image not found {img_path}")
            add_fig_caption(caption_text)

    def set_cell_shading(cell, color_hex="D9D9D9"):
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
        tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
        cell._tc.get_or_add_tcPr().append(tcMar)

    def set_table_borders(table):
        tblPr = table._tbl.tblPr
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>'
            f'<w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'<w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
            f'</w:tblBorders>'
        )
        tblPr.append(borders)

    def add_styled_table(caption, headers, rows_data, col_widths=None):
        add_tbl_caption(caption)
        tbl = doc.add_table(rows=len(rows_data)+1, cols=len(headers))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl)

        # Header
        hdr_row = tbl.rows[0]
        hdr_row._tr.get_or_add_trPr().append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        for c_idx, title in enumerate(headers):
            cell = hdr_row.cells[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_shading(cell, "D9D9D9")
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            r = p.add_run(title)
            format_run(r, size_pt=12, bold=True)

        # Data rows
        for r_idx, r_data in enumerate(rows_data):
            row = tbl.rows[r_idx+1]
            for c_idx, val in enumerate(r_data):
                cell = row.cells[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=80, bottom=80, left=140, right=140)
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                # Align center for short code/STT/HTTP methods/keys
                if c_idx == 0 or len(str(val)) <= 8 or str(val) in ['GET', 'POST', 'PUT', 'DELETE', 'INT', 'STRING', 'DATE']:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(str(val))
                format_run(r, size_pt=12)

        # Col widths
        if col_widths:
            for row in tbl.rows:
                for idx, w in enumerate(col_widths):
                    row.cells[idx].width = Inches(w)

        # Space after table
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(0)
        p_after.paragraph_format.space_after = Pt(6)
        p_after.style = 'normal'
        return tbl

    def add_toc_line(title, page_str, level=1):
        p = doc.add_paragraph()
        p.style = 'normal'
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        
        # Tabs setting with dot leader at 7.0 inches (504pt)
        pPr = p._p.get_or_add_pPr()
        tabs_xml = f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="10080"/></w:tabs>'
        pPr.append(parse_xml(tabs_xml))

        if level == 1:
            p.paragraph_format.left_indent = Inches(0)
            r1 = p.add_run(title)
            format_run(r1, size_pt=12, bold=True)
            r_tab = p.add_run("\t")
            r2 = p.add_run(page_str)
            format_run(r2, size_pt=12, bold=True)
        elif level == 2:
            p.paragraph_format.left_indent = Inches(0.2)
            r1 = p.add_run(title)
            format_run(r1, size_pt=12, bold=False)
            r_tab = p.add_run("\t")
            r2 = p.add_run(page_str)
            format_run(r2, size_pt=12, bold=False)
        else:
            p.paragraph_format.left_indent = Inches(0.4)
            r1 = p.add_run(title)
            format_run(r1, size_pt=11, bold=False)
            r_tab = p.add_run("\t")
            r2 = p.add_run(page_str)
            format_run(r2, size_pt=11, bold=False)

    print("Helper functions ready.")
    return doc, {
        'add_title': add_title,
        'add_h1': add_h1,
        'add_h2': add_h2,
        'add_h3': add_h3,
        'add_p': add_p,
        'add_bullet': add_bullet,
        'add_fig': add_fig,
        'add_fig_caption': add_fig_caption,
        'add_tbl_caption': add_tbl_caption,
        'add_styled_table': add_styled_table,
        'add_toc_line': add_toc_line
    }

if __name__ == '__main__':
    doc, h = build_report()
    print("Report initialization test successful")
